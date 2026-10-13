# Author: Andrew Fisher. Independent Decimal check against the original generator rate card.
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from datetime import date, timedelta
import hashlib, json, os, re
import openpyxl

root=Path(__file__).resolve().parent
private=Path(os.environ.get('AUDIT_PRIVATE','/workspace/private-gc500-unit1006'))
card=Path(os.environ.get('RATE_CARD','/workspace/fish/03_GC500_Delivery_Control/v7.46_gn20_install_LIVE/sources/Rate_Card_2026_1.xlsx'))
workbook=openpyxl.load_workbook(card,data_only=True)
sheet=workbook['Street Rate Card 2026']
rates={int(re.search(r'\d+',str(sheet.cell(r,1).value))[0]):(Decimal(str(sheet.cell(r,3).value)),Decimal(str(sheet.cell(r,4).value))) for r in range(66,80)}
rounding=lambda n:n.quantize(Decimal('.01'),rounding=ROUND_HALF_UP)
checks=[]
def check(name,passed):
    checks.append({'name':name,'pass':bool(passed)})

for label in ['public','candidate']:
    source=private/f'standdown-{label}-detail.json'
    if not source.exists():continue
    data=json.loads(source.read_text()); totals={k:Decimal(0) for k in ['actual','remaining','total','standdown','standdownActual','standdownRemaining']}
    for row in data['rows']:
        kva=row.get('pricedKva')
        check(f'{label} {row["id"]}: original card is the rate source',row['rateState']=='card' and kva in rates)
        if kva not in rates:continue
        daily,standby=rates[kva]
        check(f'{label} {row["id"]}: full and standdown rate match original card',abs(Decimal(str(row['dailyRate']))-daily)<Decimal('.000000001') and abs(Decimal(str(row['standdownRate']))-standby)<Decimal('.000000001'))
        start=date.fromisoformat(row['deliveredDate'] if row['actualIsConfirmed'] else row['plannedDate']); end=date(2026,10,25)
        days=[]
        while start<=end:days.append(start);start+=timedelta(days=1)
        sd=sum(d<date(2026,10,23) for d in days); race=len(days)-sd
        amount=standdown=Decimal(0)
        for contract_id in row['contractIds']:
            contract=next(c for c in data['contracts'] if c['id']==contract_id)
            qty=Decimal(str(contract['source']['quantity']))
            price=rounding(qty*(standby*sd+daily*race));down=rounding(qty*standby*sd)
            amount+=price;standdown+=down
            check(f'{label} {contract_id}: independent per-line Decimal calculation',Decimal(str(contract['charge']['forecast']))==price)
        check(f'{label} {row["id"]}: inclusive days match',row['standdownDays']==sd and row['eventDays']==race)
        check(f'{label} {row["id"]}: per-line totals preserved when grouping',Decimal(str(row['forecast']))==amount and Decimal(str(row['standdownCharge']))==standdown)
        totals['total']+=amount;totals['standdown']+=standdown
        if row['actualIsConfirmed']:totals['actual']+=amount;totals['standdownActual']+=standdown
        else:totals['remaining']+=amount;totals['standdownRemaining']+=standdown
    for k,value in totals.items():check(f'{label}: independent total {k}',Decimal(str(data['totals'][k]))==value)
summary={'author':'Andrew Fisher','rateCardSHA256':hashlib.sha256(card.read_bytes()).hexdigest(),'range':'Street Rate Card 2026!A66:D79','checks':len(checks),'failures':[c for c in checks if not c['pass']],'nativeWrites':0}
(root/'evidence'/'standdown-decimal1006.json').write_text(json.dumps(summary,indent=2)+'\n')
print(json.dumps(summary))
assert not summary['failures']
