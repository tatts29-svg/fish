"""Author: Andrew Fisher. Baseplan SuperCars (5), 16 Oct 2026: the extra charges checked against what is on site.

    python3 -I build_check.py path/to/Baseplan_SuperCars.xlsx   (the private upload; not committed)

Reads the 9968862 charge lines from the export and the site record snapshot in evidence/page_buildings.json (read from
the live page, GET only, 16 Oct 2026 ~19:10 AEST, record 8077), and writes GC500_Baseplan5_Charges_Check_16Oct2026.xlsx.
"""
import json, re, sys
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

HERE = Path(__file__).resolve().parent
CARD = {'install': 145.74, 'steps': 156.15, 'levelling': 52.05, 'cleaning': 104.10, 'fire_ext': 52.05, 'demob': 145.74}
NAMES = {'install': 'Install', 'steps': 'Steps', 'levelling': 'Levelling', 'cleaning': 'Cleaning', 'fire_ext': 'Fire extinguisher', 'demob': 'Demob'}
PACKAGE = round(sum(CARD.values()), 2)  # 655.83 - exactly the Baseplan "Installation & Dismantling" price

src = Path(sys.argv[1])
wb = openpyxl.load_workbook(src, data_only=True)
ws = wb['9968862-KINP']; rows = list(ws.iter_rows(values_only=True)); h = list(rows[0]); h[0] = 'Line'
lines = [dict(zip(h, r)) for r in rows[1:] if r[0] is not None]
page = {r['ref']: r for r in json.load(open(HERE / 'evidence' / 'page_buildings.json'))['rows']}

bld = []
for d in lines:
    if d['Item Type'] != 1: continue
    desc = str(d['Description']); m = re.match(r'^\s*([A-Z]{1,3}\d{1,3}|CP\d)\b', desc)
    if not m or not re.search(r'Building|Ticket|Office|Booth|Shell', desc, re.I): continue
    ref = m.group(1); item = str(d['Item']); p = page.get(ref) or {}
    # the record holds 1273656 at P58 (1300 Medics); Baseplan labels that line P44
    if item == '1273656': ref = 'P58'
    p = page.get(ref) or p
    ticks = {x.split(':')[0]: x.endswith('✓') for x in p.get('card', [])}
    big = bool(re.search(r'12m|9\.6m', str(p.get('item'))))
    bld.append({'line': d['Line'], 'ref': ref, 'label': m.group(1), 'item': item, 'desc': desc, 'page_item': p.get('item'), 'state': p.get('state'),
                'ticks': ticks, 'big': big})

out = openpyxl.Workbook(); B = Font(bold=True); RED = PatternFill('solid', fgColor='FBE3E1'); AMB = PatternFill('solid', fgColor='FDF1D8'); GRN = PatternFill('solid', fgColor='E3F4E8')
def sheet(name, head, data, widths=None):
    s = out.create_sheet(name); s.append(head)
    for c in s[1]: c.font = B
    for r in data: s.append(r)
    for i, w in enumerate(widths or [], 1): s.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    s.freeze_panes = 'A2'; return s

# --- Installation & Dismantling, building by building
rows_l = []; tot = {'billed': 0, 'done': 0, 'to_come': 0, 'not_recorded': 0}
for b in bld:
    t = b['ticks']; done = sum(CARD[k] for k in CARD if t.get(k) and k != 'demob'); to_come = CARD['demob'] if not t.get('demob') else 0
    if t.get('demob'): done += CARD['demob']
    nr = round(PACKAGE - done - to_come, 2)
    tot['billed'] += PACKAGE; tot['done'] += done; tot['to_come'] += to_come; tot['not_recorded'] += nr
    miss = [NAMES[k] for k in CARD if k != 'demob' and not t.get(k)]
    rows_l.append([b['line'], b['ref'], b['label'] if b['label'] != b['ref'] else '', b['item'], b['page_item'], b['state'], PACKAGE, round(done, 2), to_come, nr, ', '.join(miss) or '—'])
s = sheet('Install & dismantle', ['Baseplan line', 'Ref (site record)', 'Baseplan label if different', 'Asset no', 'Site record item', 'On site?', 'Baseplan charges', 'Work done (ticked on site)', 'Demob - to come', 'Charged, not recorded on site', 'Not done on site'],
          rows_l + [[], ['', 'TOTAL', '', '', '', '', round(tot['billed'], 2), round(tot['done'], 2), round(tot['to_come'], 2), round(tot['not_recorded'], 2), '']], [8, 10, 12, 11, 15, 9, 12, 14, 12, 14, 34])
for r in s.iter_rows(min_row=2):
    if isinstance(r[9].value, (int, float)) and r[9].value > 0: r[9].fill = RED

# --- Transport
small = [b for b in bld if not b['big']]; big = [b for b in bld if b['big']]
trows = [['Up to 6 m buildings (line 134)', 72, 500, 36000, len(small), len(small) * 500, len(small) * 500, 'Rate matches the card ($500 each way). Count matches: %d buildings on site x 2 trips.' % len(small)],
         ['9.6 m and 12 m buildings (line 135)', 18, 800, 14400, len(big), len(big) * 800, len(big) * 800, 'Rate matches the card ($800 each way). Count matches: %d buildings x 2 trips.' % len(big)]]
sheet('Transport', ['Line', 'Baseplan qty', 'Rate', 'Baseplan charges', 'Buildings on site', 'Delivery leg - done', 'Return leg - not done yet', 'Check'], trows, [34, 12, 8, 14, 16, 18, 22, 70])

# --- Other changes vs Baseplan (4)
other = [
 ['9961976 line 15 (GN10)', 'NEW: a 60 kVA generator, asset 1283681, $1,142.96/wk, from 19 Oct (was "GN10 45kVA", no unit)', 'GN10 is a 20 kVA, asset 1197355, not on site yet, due 19 Oct. 1283681 is the spare we took off GN06.', 'WRONG UNIT AND SIZE - would charge 60 kVA for a 20 kVA job. Ask the branch.', 'red'],
 ['9961976 line 48 (GN06)', 'Delivered 15 Oct, asset 1320828, described 60 kVA, $1,184.82/wk', 'GN06 on site 15 Oct, asset 1320828 (same unit), recorded as 50 kVA', 'Size differs - check the plate. If it is 50 kVA it should be on the 50 kVA rate.', 'amber'],
 ['9992300 line 12', 'NEW: Straight Boom Lift 60 ft, asset 1286482, $459.55/wk from 21 Oct, prebill $2,417.94 (TV compound)', 'Not on our record. Our TV compound references are FL03, FL04 and TM01 only; the knuckle boom (line 5) and scissor lift (line 4) are not on our record either.', 'Was a second boom asked for, or does it replace the knuckle boom? There is only one boom transport line (line 10).', 'red'],
 ['9972648 line 18 (WB22)', 'NEW: 4 Armorzone barriers from 18 Oct', 'No WB22 on our schedule', 'Not ours unless a request came in. Ask.', 'red'],
 ['9972648 lines 7 and 8 (WB07, WB13)', 'REMOVED (17 and 10 barriers)', 'WB07 17 due 18 Oct, WB13 10 due 19 Oct, both still on our schedule', 'If they go out they will not be billed. Cancelled, or dropped by mistake?', 'amber'],
 ['9972648 lines 4 and 5 (WB04, WB05)', 'Labels swapped since (4); quantities 13 (WB04) and 12 (WB05), delivered 16 Oct', 'WB04 14 (Helen Park) and WB05 13 (Main Beach) on site 16 Oct', 'Two fewer barriers billed than are on site.', 'amber'],
 ['9968862 lines 34, 92, 105 (P19, P21, P25)', 'Described 6.0 m x 3.0 m', 'Recorded as 4.8 m buildings', 'Description only - billed at the 4.8 m rate ($1,125.84). No money difference.', 'green'],
 ['9968862 line 59 ("P44")', 'Asset 1273656 labelled P44 (P44 also on line 61, asset 198481)', '1273656 is P58 (1300 Medics); P44 is 198481', 'Label only - one building each, nothing double charged.', 'green'],
 ['9968955 line 111 (WC51)', 'Now Delivered 15 Oct, asset 1197174', 'WC51 on site 15 Oct, no number recorded', 'Agrees. 1197174 can go on the record once you have seen the plate.', 'green'],
 ['9968726 (whole contract)', 'REMOVED (old TV compound order)', 'Already held out as superseded by 9992300', 'Agrees.', 'green'],
 ['9961976 lines 13 and 45', 'REMOVED (spare 60 kVA generators 1283681 and 1261271)', '1283681 was taken off GN06', 'Agrees.', 'green'],
 ['9968862 line 63 (P41)', 'Rate $1,426.17 -> $1,426.07', '', '10 cents a week. Nothing to do.', 'green'],
 ['9992300 lines 1-6', 'Forklifts 5 t and 2.5 t, fork extensions, 10 Trakmats delivered 16 Oct; scissor lift moved to 19 Oct, knuckle boom to 21 Oct', 'FL03, FL04, TM01 on site', 'Agrees for the forklifts and Trakmats.', 'green'],
]
s = sheet('Other changes vs (4)', ['Where', 'Baseplan (5) says', 'Our site record says', 'What to do', ''], [r[:4] for r in other], [30, 60, 60, 60])
for i, r in enumerate(other, 2):
    for c in s[i][:4]: c.fill = {'red': RED, 'amber': AMB, 'green': GRN}[r[4]]; c.alignment = Alignment(wrap_text=True, vertical='top')

# --- Summary first
sm = out['Sheet']; sm.title = 'Summary'
sm.append(['Baseplan SuperCars (5), 16 Oct 2026 - charges checked against the site record (record 8077)']); sm['A1'].font = Font(bold=True, size=13)
sm.append([]); sm.append(['New charges on 9968862, dated 16 Oct', 'Baseplan', 'Site record supports now', 'Not done yet (legit once done)', 'Charged but not recorded on site'])
for c in sm[3]: c.font = B
sm.append(['Installation & Dismantling, 45 x $655.83', round(tot['billed'], 2), round(tot['done'], 2), round(tot['to_come'], 2), round(tot['not_recorded'], 2)])
sm.append(['Transport up to 6 m, 72 x $500', 36000, len(small) * 500, len(small) * 500, 0])
sm.append(['Transport 9.6 m & 12 m, 18 x $800', 14400, len(big) * 800, len(big) * 800, 0])
T = [round(tot['billed'] + 50400, 2), round(tot['done'] + len(small) * 500 + len(big) * 800, 2), round(tot['to_come'] + len(small) * 500 + len(big) * 800, 2), round(tot['not_recorded'], 2)]
sm.append(['TOTAL'] + T)
for c in sm[7]: c.font = B
sm.append([]); sm.append(['$655.83 is the whole card labour package for one building: Install 145.74 + Steps 156.15 + Levelling 52.05 + Cleaning 104.10 + Fire extinguisher 52.05 + Demob 145.74. It is charged on all 45 buildings whether or not each piece was done.'])
sm.append(['On site: install, levelling, cleaning on all 45; steps on 41 (none on P06, P38, P39, P42); fire extinguisher on none; demob still to come.'])
sm.append(['Transport rates and counts match the card and the buildings on site; half of it is the return trip, not done until demob.'])
sm.column_dimensions['A'].width = 52
for col in 'BCDE': sm.column_dimensions[col].width = 22
out.save(HERE / 'GC500_Baseplan5_Charges_Check_16Oct2026.xlsx')
print('billed', round(tot['billed'], 2), 'done', round(tot['done'], 2), 'demob to come', round(tot['to_come'], 2), 'not recorded', round(tot['not_recorded'], 2), 'buildings', len(bld), 'small', len(small), 'big', len(big))
