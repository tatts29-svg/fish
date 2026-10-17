"""Author: Andrew Fisher. Buildings - reference and all charges from the site record (16 Oct 2026).

    python3 -I build_xlsx.py buildings.json out.xlsx

Charges are counted only where the record has them: a labour line when it is ticked, a transport leg when its tick is
on. Hire is the contract amount the page carries for the building. Baseplan (5)'s 9968862 install & dismantle
($655.83 a building) and transport ($500 / $800 each way) are shown beside it so the differences stand out. Ex GST.
"""
import json, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

rows = json.load(open(sys.argv[1]))
LAB = [('install', 'Install'), ('steps', 'Steps'), ('levelling', 'Levelling'), ('cleaning', 'Cleaning'), ('fire_ext', 'Fire extinguisher'), ('demob', 'Demob')]
r2 = lambda v: round(v + 1e-9, 2)
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Buildings'
B = Font(bold=True); RED = PatternFill('solid', fgColor='FBE3E1'); GRY = PatternFill('solid', fgColor='EEEEEE')
head = ['Ref', 'Name', 'Building', 'Asset no', 'Contract / line', 'On site', 'Hire (contract)'] + [n for _, n in LAB] + ['Delivery', 'Pick-up', 'Labour + transport', 'TOTAL (my record)', 'Baseplan install & dismantle + transport', 'Baseplan over my record']
ws.append(head)
for c in ws[1]: c.font = B; c.alignment = Alignment(wrap_text=True, vertical='top')
on = [r for r in rows if r['state'] == 'on site' and not r['cancelled'] and not r['moved']]
off = [r for r in rows if r not in on]
tot = [0.0] * (len(head))
def num(v): return v if isinstance(v, (int, float)) else None
for r in sorted(on, key=lambda x: (x['ref'][0], int(''.join(ch for ch in x['ref'] if ch.isdigit()) or 0), x['ref'])):
    lab = [r2(r['lab'][k]['rate']) if k in r['lab'] and r['lab'][k]['ticked'] and r['lab'][k]['rate'] else None for k, _ in LAB]
    dl = r['tr'].get('delivery') or {}; pk = r['tr'].get('pickup') or {}
    dv = num(dl.get('charge')) if dl.get('tick') == 'on' else None; pv = num(pk.get('charge')) if pk.get('tick') == 'on' else None
    lt = r2(sum(x or 0 for x in lab) + (dv or 0) + (pv or 0)); total = r2(lt + (r['hire'] or 0))
    bp = None
    if r['contract'] == '9968862':
        big = '12m' in str(r['item']) or '9.6m' in str(r['item'])
        bp = r2(655.83 + (1600 if big else 1000))
    over = r2(bp - lt) if bp is not None else None
    row = [r['ref'], r['name'], r['item'], r['assetNo'], (r['contract'] + ' / ' + str(r['line'])) if r['contract'] else '—', 'Yes', r['hire']] + lab + [dv, pv, lt, total, bp, over]
    ws.append(row)
    for i, v in enumerate(row):
        if isinstance(v, (int, float)) and i >= 6: tot[i] += v
    if over and over > 0.005: ws.cell(ws.max_row, len(head)).fill = RED
ws.append([])
ws.append(['TOTAL', '%d buildings on site' % len(on), '', '', '', ''] + [r2(t) if i >= 6 else None for i, t in enumerate(tot)][6:])
for c in ws[ws.max_row]: c.font = B
for i, w in enumerate([7, 30, 13, 10, 15, 8, 12, 9, 9, 10, 10, 11, 9, 10, 10, 13, 14, 16, 14], 1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ws.freeze_panes = 'C2'
for row in ws.iter_rows(min_row=2):
    for c in row[6:]:
        if isinstance(c.value, (int, float)): c.number_format = '$#,##0.00'
s2 = wb.create_sheet('Not on site')
s2.append(['Ref', 'Name', 'Building', 'Record says']);
for c in s2[1]: c.font = B
for r in off: s2.append([r['ref'], r['name'], r['item'], 'cancelled' if r['cancelled'] else 'moved away' if r['moved'] else (r['state'] or 'no record')])
s2.column_dimensions['B'].width = 34
wb.save(sys.argv[2])
labels = head[6:]
print({labels[i - 6]: r2(tot[i]) for i in range(6, len(head))}, len(on), 'on site,', len(off), 'not')
