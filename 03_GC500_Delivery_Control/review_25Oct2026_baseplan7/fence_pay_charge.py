# Author: Andrew Fisher. 25 Oct 2026 - the fencing behind the $405,150.96: what Coates charges the V8s and what Coates pays
# Advanced, item by item, week by week and docket by docket, from the live page's own pricing (dump_fence_pay_charge.cjs and
# the page's paid split). Andrew: "can you give me breakdown of what are costs are we are charging to get to thisd please.
# what we pay and then what we charge so i can show gow we got this figure $405,150.96". Reads private dumps only; writes
# GC500_Fencing_Charge_and_Cost_25Oct2026.xlsx and evidence/fence_pay_charge.json (no names, no links). Then, Andrew: "just want
# to make sure what ever we have receievd for from their invoices have matched correctly ... what we charge the v8s but also
# know via break down of what we pay advance based off everything we have" - so every Advanced order and invoice is set
# against the dockets and notes on our record, with what is not on any of their sheets named. Then, Andrew: "we charge once" -
# live v12.16 (14:18) takes the fence under the scrim off the clean charge, docket by docket - and "show me it like this item
# the total then rate we pay advnanced then the rate we charge v8s then the total": section 1 is in that order.
#   python3 fence_pay_charge.py <fence_pay_charge.json> <fence_split.json> <fence_invoices.json>
import collections, json, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

D, SPLIT, INV = json.load(open(sys.argv[1])), json.load(open(sys.argv[2])), json.load(open(sys.argv[3]))
HERE = Path(__file__).resolve().parent
cents = lambda v: int(round((v or 0) * 100))
d = lambda c: c / 100
COLS = {c['key']: c for c in D['columns']}
use = [x for x in D['dockets'] if x['usable']]
ITEMS = ['clean', 'scrim', 'ccb_demarc', 'ccb_event', 'flat_feet', 'labour', 'relocation', 'ped_gates', 'fence_blocks', 'v_gates', 'bases']
NAME = {'clean': 'Clean fence (not scrimmed)', 'scrim': 'Scrim fence (braced, shade cloth on)', 'ccb_demarc': 'CCB demarcation', 'ccb_event': 'CCB event',
        'flat_feet': 'CCB flat feet', 'labour': 'Labour (docket 36511)', 'relocation': 'Relocation', 'ped_gates': 'Pedestrian gates',
        'fence_blocks': 'Fence blocks', 'v_gates': 'Vehicle gates', 'bases': 'Bases'}
PAYWHY = {'clean': "Advanced's rate sheet", 'scrim': "Advanced's rate sheet (braced for scrim)", 'ccb_demarc': "Advanced's rate sheet", 'ccb_event': "Advanced's rate sheet",
          'flat_feet': 'typed on our record 15 Oct (no flat-feet line on their sheet)', 'labour': "Advanced's rate sheet", 'relocation': 'they bill relocation by the hour, $100',
          'ped_gates': "typed on our record 15 Oct; their invoices say $61 incl one wheel", 'fence_blocks': 'no per-block line on their sheet: the cost is on their invoice',
          'v_gates': "Advanced's rate sheet", 'bases': "on their invoices"}

def charge_from(k):
    c = COLS.get(k) or {}
    if k == 'v_gates': return 'included in the metre rate (the project manager, 1 Oct 2026)'
    if c.get('card_line') and not c.get('settled_by'): return '2026 street card: ' + c['card_line']
    if c.get('card_line'): return '2026 street card: ' + c['card_line'] + ' (confirmed ' + c['settled_on'] + ')'
    if c.get('settled_on'): return 'your rate, ' + c['settled_on']
    return ''

# --- the charge side: our docket quantities at our rates, each docket line rounded to the cent as the page prices it
charge = collections.OrderedDict((k, {'qty': 0, 'c': 0, 'rate': None}) for k in ITEMS)
pay = collections.OrderedDict((k, {'qty': 0, 'c': 0, 'rates': set()}) for k in ITEMS)
for x in use:
    for l in x['lines']:
        a = charge[l['column']]; a['qty'] += l['qty']; a['c'] += cents(l['cost']); a['rate'] = l['rate']
    # the paid side: Advanced's approved billed lines where a purchase order carries them, else our quantities at their sheet
    for l in (x['supplier']['lines'] if x['supplier'] else x['lines']):
        p = pay[l['column']]; p['qty'] += l['qty']; p['c'] += l['cents'] if x['supplier'] else cents(l['paid'])
        r = l['rate'] if x['supplier'] else l['paid_rate']
        if r is not None: p['rates'].add(r)
TC = sum(a['c'] for a in charge.values()); TP = sum(p['c'] for p in pay.values())
assert TC == cents(sum(x['cost_total'] for x in use)), TC
assert TP == cents(sum(x['paid_total'] for x in use)), TP

S = SPLIT['split']; O = SPLIT['orders']
EX = SPLIT['extras']
tl = [x for x in EX if x['column'] == 'team_leader']; mats = [x for x in EX if x['column'] != 'team_leader']
TL, MAT = sum(x['cents'] for x in tl), sum(x['cents'] for x in mats)
ADJ = cents(O['adjustment']); GREEN = cents(S['green'])
FENCE_COST = TP + TL + MAT + ADJ
assert FENCE_COST == cents(S['paid']), (FENCE_COST, S['paid'])
ALL_IN = FENCE_COST + GREEN

# --- the scrim check: Advanced's invoiced clean = the paper's clean less their scrim and 2.5 m a vehicle gate
q = lambda lines, col: sum(l['qty'] for l in lines if l['column'] == col)
qp = lambda lines, col: sum(l['qty'] for l in lines if l['column'] == col and l['qty'] > 0)   # what the paper says, without the v12.16 line
under = lambda lines: -sum(l['qty'] for l in lines if l['column'] == 'clean' and l['qty'] < 0)
check = []
for x in use:
    if not x['supplier']: continue
    oc, ac, as_, av = qp(x['lines'], 'clean'), q(x['supplier']['lines'], 'clean'), q(x['supplier']['lines'], 'scrim'), q(x['supplier']['lines'], 'v_gates')
    check.append(round(oc - as_ - 2.5 * av - ac, 2))
exact = sum(1 for r in check if r == 0)
SCRIM_M = charge['scrim']['qty']; CLEAN_RATE = charge['clean']['rate']
ONCE = -sum(cents(l['cost']) for x in use for l in x['lines'] if l['column'] == 'clean' and l['qty'] < 0)   # what v12.16 took off the charge
PAPER_CLEAN = sum(qp(x['lines'], 'clean') for x in use)
assert round(sum(under(x['lines']) for x in use), 2) == SCRIM_M, 'every scrim metre has its line'
BASEPLAN = 30910683

# --- by week
WEEK_OF_SHEET = {'CON WK6': 'Week 6', 'CON WK5': 'Week 5', 'CON WK4': 'Week 4', 'CON WK3': 'Week 3', 'CON WK2': 'Week 2', 'CON WK1': 'Week 1', 'EVENT WEEK': 'Event Week'}
po_week = {p['number']: WEEK_OF_SHEET.get(p['programme_sheet']) for p in D['pos'] if p.get('programme_sheet')}
SPANS = [('Week 6', '2026-09-07', '2026-09-13'), ('Week 5', '2026-09-14', '2026-09-20'), ('Week 4', '2026-09-21', '2026-09-27'), ('Week 3', '2026-09-28', '2026-10-04'),
         ('Week 2', '2026-10-05', '2026-10-11'), ('Week 1', '2026-10-12', '2026-10-18'), ('Event Week', '2026-10-19', '2026-10-25')]
week_of_day = lambda day: next((w for w, a, b in SPANS if a <= (day or '') <= b), None)
W = collections.OrderedDict((w, {'dockets': 0, 'approved': 0, 'charge': 0, 'pay': 0, 'extras': 0, 'notes': 0, 'adj': 0, 'from': None, 'to': None}) for w, _, _ in SPANS)
for x in use:
    w = W[x['week']]; w['dockets'] += 1; w['approved'] += 1 if x['supplier'] else 0; w['charge'] += cents(x['cost_total']); w['pay'] += cents(x['paid_total'])
    if x['date']: w['from'] = min(w['from'] or x['date'], x['date']); w['to'] = max(w['to'] or x['date'], x['date'])
for x in EX: W[po_week[x['po']]]['extras'] += x['cents']
for n in D['notes']:
    if n.get('usable') is not False: W[week_of_day(n['date'])]['notes'] += cents(n['cost'])
for o in O['orders']: W[po_week[o['number']]]['adj'] += cents(o['adjustment'])
assert sum(w['charge'] for w in W.values()) == TC and sum(w['pay'] + w['extras'] + w['notes'] + w['adj'] for w in W.values()) == ALL_IN
POS = {po_week[p['number']]: p for p in D['pos'] if p['number'] in po_week}

# --- Advanced's orders and invoices against our record: what each sheet holds, what of ours is on none of them
ON_SHEET = {('red', r['id']) for r in INV['red'] if r['onSheet']} | {('green', n['id']) for n in INV['green'] if n['onSheet']}
off_red = [r for r in INV['red'] if not r['onSheet']]; off_green = [n for n in INV['green'] if not n['onSheet']]
ADV = [o for o in INV['pos'] if o['po'].get('supplier_name') == 'Advanced Temporary Fencing']
OTHER = [o for o in ADV if not o['weeks']]                       # an Advanced invoice tied to no programme week
OTHER_C = sum(cents(o['po']['amount']) for o in OTHER if o['po'].get('amount') is not None)
EVERYTHING = ALL_IN + OTHER_C
carried = {}                                                     # what the page carries for each work week (= the week's all-in above)
for k, w in W.items(): carried[k] = w['pay'] + w['extras'] + w['notes'] + w['adj']
def match_words(o, k):
    st, sh, po = o['state'] or {}, o['sheet'], o['po']
    week_off = [r for r in off_red if r['week'] == k]; notes_off = [n for n in off_green if week_of_day(n['date']) == k]
    words = []
    if sh and sh['valid'] and sh['difference'] == 0: words.append("Advanced's approved sheet = the order, to the cent")
    elif sh and sh['valid']: words.append("sheet $%s, %s the order; the order stands" % (format(sh['total'], ',.2f'), ('%d cents over' % round(sh['difference'] * 100)) if abs(sh['difference']) < 1 else '$' + format(abs(sh['difference']), ',.2f') + (' over' if sh['difference'] > 0 else ' under')))
    elif st.get('state') == 'receipted': words.append('invoiced; no sheet to check line by line; the order stands')
    elif st.get('state') == 'awaiting': words.append('nothing billed yet')
    if k == 'Week 6': words.append('our %d dockets at their rates come to $%s, so the invoice is $%s less' % (len(week_off), format(st.get('est') or 0, ',.2f'), format((st.get('est') or 0) - (po.get('amount') or 0), ',.2f')))
    elif week_off and k != 'Event Week': words.append('not on their sheet: docket ' + ', '.join(str(r['no']) for r in week_off) + ' ($' + format(sum(r['paid'] for r in week_off), ',.2f') + ' at their rates)')
    if notes_off and k != 'Event Week': words.append('not on their sheet: labour note ' + ', '.join(str(n['no']) for n in notes_off) + ' ($' + format(sum(n['cost'] for n in notes_off), ',.2f') + ')')
    if k == 'Event Week': words.append('%d dockets and %d labour notes on our record, at their rates' % (len(week_off), len(notes_off)))
    if sh and (sh.get('source') or {}).get('po_as_printed') and sh['source']['po_as_printed'] != po['number']: words.append('their sheet prints the order as ' + sh['source']['po_as_printed'])
    return '; '.join(words)
ORDERS = []
for k in W:
    o = next((x for x in ADV if po_week.get(x['po']['number']) == k), None)
    if not o: continue
    sh = o['sheet']; po = o['po']
    ORDERS.append([k, po['number'], sh['total'] if sh else None, po.get('amount'), po.get('invoice_no'), match_words(o, k), d(carried[k])])
dm = next((x for x in ADV if po_week.get(x['po']['number']) is None and x['weeks'] and all('Demob' in w for w in x['weeks'])), None)
if dm: ORDERS.append(['Demob', dm['po']['number'], None, None, None, 'nothing billed yet; no papers yet', None])
for o in OTHER: ORDERS.append(['Other', o['po']['number'], None, o['po'].get('amount'), o['po'].get('invoice_no'), (o['po'].get('note') or '') + '; tied to no docket (the collection forms carry counts, no money); not in the page\'s fencing cost', o['po'].get('amount')])
assert sum(cents(r[6]) for r in ORDERS if r[6] is not None) == EVERYTHING

# ------------------------------------------------------------------------------------------------- workbook
wb = Workbook(); wb.remove(wb.active)
H = Font(bold=True); BIG = Font(bold=True, size=13); AMB = PatternFill('solid', fgColor='FFF2CC'); GRN = PatternFill('solid', fgColor='E2EFDA'); BLU = PatternFill('solid', fgColor='DDEBF7'); RED = PatternFill('solid', fgColor='F8CBAD')
M = '$#,##0.00'; Q = '#,##0.0#'
def widths(s, ws):
    for i, w in enumerate(ws): s.column_dimensions[chr(65 + i) if i < 26 else 'A' + chr(65 + i - 26)].width = w
def put(s, row, vals, font=None, fill=None, fmt=None):
    s.append(list(vals)); r = s.max_row
    for j, c in enumerate(s[r]):
        c.alignment = Alignment(wrap_text=True, vertical='top')
        if font: c.font = font
        if fill: c.fill = fill
        if fmt and j in fmt: c.number_format = fmt[j]
    return r

s = wb.create_sheet('Summary'); widths(s, [34, 12, 7, 15, 15, 15, 44, 44])
put(s, 0, ['GC500 fencing (9990776): what we charge the V8s and what we pay Advanced - signed dockets 9 Sep to 23 Oct 2026, ex GST'], BIG)
put(s, 0, ['From the live record on 25 Oct 2026 (page v12.16), priced by the page itself: 167 signed red-book dockets, the fence under the scrim charged once. Our rates are the 2026 street card or your named rates. '
           "Advanced's costs are their approved pricing on the purchase orders for weeks 1-5 (115 dockets), and their rate sheet for week 6 and the event week (52 dockets) until their invoices land."])
put(s, 0, [])
put(s, 0, ['1. What we charge the V8s, item by item'], H, BLU)
put(s, 0, ['Item', 'Total', 'Unit', 'Rate we pay Advanced', 'Rate we charge V8s', 'Total we charge', 'Our rate comes from', "Advanced's rate comes from"], H)
for k in ITEMS:
    a = charge[k]
    if not a['qty'] and k != 'bases': continue
    if k == 'bases': continue
    pr = sorted(pay[k]['rates']); unit = COLS.get(k, {}).get('unit') or ''
    put(s, 0, [NAME[k], a['qty'], unit, pr[0] if len(pr) == 1 and k != 'relocation' else None, a['rate'], d(a['c']), charge_from(k), PAYWHY[k]],   # relocation: their rate is per hour, not per metre
        fmt={1: Q, 3: '$#,##0.00', 4: '$#,##0.00##', 5: M})
put(s, 0, ['Total we charge', None, None, None, None, d(TC)], H, GRN, fmt={5: M})
put(s, 0, ["Clean fence is the papers' %s m less the %s m of fence under the scrim, which is charged once, in the scrim (v12.16). Each docket line is rounded to the cent, as the page prices it." % (format(PAPER_CLEAN, ',.1f'), format(SCRIM_M, ',.1f'))])
put(s, 0, [])
put(s, 0, ['2. What we pay Advanced for the same work'], H, BLU)
put(s, 0, ['Item', "Qty Advanced bill", 'Unit', 'Their rate', 'Note', 'We pay'], H)
for k in ITEMS:
    p = pay[k]
    if not p['qty']: continue
    pr = sorted(p['rates']); unit = {'relocation': 'h'}.get(k, COLS.get(k, {}).get('unit') or 'each')
    note = {'clean': "their clean metres leave out the scrimmed fence and the gate openings (see 5)", 'ccb_demarc': 'on weeks 1-5 Advanced bill all CCB on this one line, same $5.40',
            'v_gates': 'billed by Advanced; not charged to the V8s (inside the metre rate)', 'bases': 'on their week 3 invoice', 'relocation': 'relocation billed as labour hours'}.get(k, '')
    put(s, 0, [NAME[k], p['qty'], unit, pr[0] if len(pr) == 1 else None, note, d(p['c'])], fmt={1: Q, 3: '$#,##0.00', 5: M})
put(s, 0, ['Fence, barriers and gates on the 167 dockets', None, None, None, None, d(TP)], H, fmt={5: M})
put(s, 0, ['Team leader', sum(x['qty'] for x in tl), 'days', 200, 'on the purchase orders, weeks 1-5', d(TL)], fmt={1: Q, 3: '$#,##0.00', 5: M})
for x in mats: put(s, 0, [x['name'], x['qty'], x['unit'], x['rate'], 'on purchase order ' + x['po'], d(x['cents'])], fmt={1: Q, 3: '$#,##0.00', 5: M})
for o in O['orders']:
    if o['adjustment']: put(s, 0, ['Invoice under the sheet rates: ' + o['number'] + ' (' + po_week[o['number']] + ')', None, None, None,
                                   'receipt ' + o['receipt'] + ' $' + format(o['amount'], ',.2f') + ' against $' + format(o['counted'], ',.2f') + ' at the sheet rates; the invoice stands (the project manager, 23 Oct)', o['adjustment']], fmt={5: M})
put(s, 0, ['Fencing cost on the page', None, None, None, None, d(FENCE_COST)], H, fmt={5: M})
put(s, 0, ['Labour notes (green book)', S['green_hours'], 'h', S['green_rate'], str(S['green_notes']) + " service notes: Advanced's labour, a cost only (inside our metre rates)", d(GREEN)], fmt={1: Q, 3: '$#,##0.00', 5: M})
put(s, 0, ['We pay for the work on the 167 dockets', None, None, None, None, d(ALL_IN)], H, fmt={5: M})
for o in OTHER: put(s, 0, ['Other Advanced invoice: order ' + o['po']['number'], None, None, None, (o['po'].get('note') or '') + ' (receipt ' + str(o['po'].get('invoice_no')) + '); not in the page\'s fencing cost', o['po'].get('amount')], fmt={5: M})
put(s, 0, ['Total we pay Advanced, everything to date', None, None, None, None, d(EVERYTHING)], H, GRN, fmt={5: M})
put(s, 0, [])
put(s, 0, ["3. What Advanced have billed, order by order - and whether it matches our record"], H, BLU)
put(s, 0, ['Week', 'Order', "Advanced's sheet", 'Order amount', 'Receipt', 'How it matches', 'We carry'], H)
for r in ORDERS: put(s, 0, r, fmt={2: M, 3: M, 6: M})
put(s, 0, ['Total', None, None, None, None, None, d(EVERYTHING)], H, GRN, fmt={6: M})
put(s, 0, [])
put(s, 0, ['4. The difference'], H, BLU)
put(s, 0, ['We charge the V8s', None, None, None, None, d(TC)], fmt={5: M})
put(s, 0, ['We pay Advanced', None, None, None, None, d(EVERYTHING)], fmt={5: M})
put(s, 0, ['Margin', None, None, None, '%.1f%% of what we charge' % (100 * (TC - EVERYTHING) / TC), d(TC - EVERYTHING)], H, fmt={5: M})
put(s, 0, [])
put(s, 0, ['5. The scrim fence is charged once (Andrew, 25 Oct 2026: "we charge once")'], H, GRN)
for line in [
    "The scrim rate ($25.53 a metre, set 24 Sep) is the fence with shade cloth on it: card fence $16.16 + card shade cloth $9.37. "
    "Advanced bill braced-for-scrim the same way, all in ($24.40).",
    "The crews' papers write the whole fence run as clean, scrimmed sections included. On %d of the %d dockets Advanced have invoiced, their clean metres are "
    "exactly the paper's clean less the scrim and 2.5 m for each vehicle gate. Docket 36532: paper 150 m clean and 132.5 m scrim; Advanced bill 12.5 m clean, "
    "132.5 m scrim and 2 vehicle gates." % (exact, len(check)),
    "Since v12.16 (live 25 Oct 14:18) each docket with scrim carries one more line: its scrim metres off the clean at $16.1607. That took $%s off the charge "
    "(was $%s); on the paid side it took the event week's estimate down $6,100, the way Advanced bill it." % (format(d(ONCE), ',.2f'), format(d(TC + ONCE), ',.2f')),
    "Against Baseplan's $309,106.83 for 9990776, our signed dockets come to $%s: Baseplan is short $%s." % (format(d(TC), ',.2f'), format(d(TC - BASEPLAN), ',.2f'))]:
    put(s, 0, [line])
put(s, 0, [])
put(s, 0, ['6. Also worth knowing'], H, BLU)
for line in [
    "Pedestrian gates: we pay $61.00 each and charge the card's $36.18, so each one loses $24.82 (22 gates, $546.04 under).",
    "Vehicle gates: Advanced bill $50 each (%d, $%s); we charge none on their own, as the project manager settled (inside the metre rate)." % (pay['v_gates']['qty'], format(d(pay['v_gates']['c']), ',.2f')),
    "The event week (19-23 Oct, 45 dockets) is costed at Advanced's rate sheet until their invoice comes in (order 4647522 has no amount yet), so the paid side can still move.",
    "Not in these figures: the demob (removal is inside the metre rate; its papers are still to come) and the fence team's event labour under the C004 scope "
    "(3.15 Security & Demarcation Fencing, 213 h at $91.14 = $19,412.82), which is a separate charge."]:
    put(s, 0, [line])

s = wb.create_sheet('By week'); widths(s, [11, 22, 9, 10, 15, 15, 13, 13, 13, 15, 12, 13, 11])
put(s, 0, ['Week', 'Dockets dated', 'Dockets', "On Advanced's invoice pricing", 'We charge', 'We pay: dockets', 'We pay: team leader + materials', 'We pay: labour notes',
           'Invoice adjustment', 'We pay: all in', 'Purchase order', 'Order amount', 'Receipt'], H)
for k, w in W.items():
    p = POS.get(k) or {}
    put(s, 0, [k, (w['from'] or '') + (' to ' + w['to'] if w['to'] and w['to'] != w['from'] else ''), w['dockets'], w['approved'], d(w['charge']), d(w['pay']), d(w['extras']), d(w['notes']), d(w['adj']),
               d(w['pay'] + w['extras'] + w['notes'] + w['adj']), p.get('number'), p.get('amount'), p.get('invoice_no')], fmt={4: M, 5: M, 6: M, 7: M, 8: M, 9: M, 11: M})
put(s, 0, ['Total', None, sum(w['dockets'] for w in W.values()), sum(w['approved'] for w in W.values()), d(TC), d(TP), d(TL + MAT), d(GREEN), d(ADJ), d(ALL_IN)], H, GRN, fmt={4: M, 5: M, 6: M, 7: M, 8: M, 9: M})
for o in OTHER: put(s, 0, ['Other', 'order ' + o['po']['number'], None, None, None, None, None, None, None, o['po'].get('amount'), o['po']['number'], o['po'].get('amount'), o['po'].get('invoice_no')], fmt={9: M, 11: M})
put(s, 0, ['Total, everything', None, None, None, d(TC), None, None, None, None, d(EVERYTHING)], H, GRN, fmt={4: M, 9: M})
put(s, 0, ['Week 6 includes one undated row from the fencing cost tracker (232.5 m clean, 3 vehicle gates). Labour notes go to the week of their date; extras and invoice adjustments to their order\'s week.'])
s.freeze_panes = 'A2'

s = wb.create_sheet('Dockets'); widths(s, [9, 11, 11, 30, 9, 9, 9, 9, 9, 8, 8, 9, 8, 8, 12, 12, 30, 10, 10, 9, 12, 16])
head = ['Docket', 'Date', 'Week', 'Location', 'Clean m', 'Scrim m', 'CCB demarc m', 'CCB event m', 'Flat feet m', 'Ped gates', 'V gates', 'Relocation m', 'Blocks', 'Labour h',
        'We charge', 'We pay', 'Paid at', 'Advanced clean m', 'Advanced scrim m', 'Advanced V gates', 'Fence under the scrim, off the clean (m)', "On an Advanced sheet"]
put(s, 0, head, H)
for x in sorted(use, key=lambda x: (x['date'] or '', str(x['no'] or ''))):
    qq = lambda col: qp(x['lines'], col) or None
    sl = x['supplier']['lines'] if x['supplier'] else None
    adv = (q(sl, 'clean'), q(sl, 'scrim'), q(sl, 'v_gates')) if sl else (None, None, None)
    twice = under(x['lines']) or None
    r = put(s, 0, [x['no'] or 'tracker row', x['date'], x['week'], x['location'], qq('clean'), qq('scrim'), qq('ccb_demarc'), qq('ccb_event'), qq('flat_feet'), qq('ped_gates'), qq('v_gates'),
                   qq('relocation'), qq('fence_blocks'), qq('labour'), x['cost_total'], x['paid_total'],
                   ("Advanced's invoice pricing, order " + x['supplier']['po']) if x['supplier'] else "Advanced's rate sheet", adv[0], adv[1], adv[2], twice,
                   'yes' if ('red', x['id']) in ON_SHEET else ('event week: not billed yet' if x['week'] == 'Event Week' else 'week 6: no sheet' if x['week'] == 'Week 6' else 'NO')], fmt={14: M, 15: M})
    if twice:
        for c in s[r]: c.fill = AMB
put(s, 0, ['Total', None, None, None, PAPER_CLEAN, charge['scrim']['qty'], charge['ccb_demarc']['qty'], charge['ccb_event']['qty'], charge['flat_feet']['qty'], charge['ped_gates']['qty'],
           charge['v_gates']['qty'], charge['relocation']['qty'], charge['fence_blocks']['qty'], charge['labour']['qty'], d(TC), d(TP)], H, GRN, fmt={14: M, 15: M})
s.freeze_panes = 'E2'

s = wb.create_sheet('Labour notes and extras'); widths(s, [24, 12, 30, 10, 8, 10, 12])
put(s, 0, ['Labour notes (green book): Advanced Environmental Services service notes, a cost only'], H, BLU)
put(s, 0, ['Note', 'Date', 'Hours', 'Rate', 'On an Advanced sheet', None, 'We pay'], H)
for n in sorted(D['notes'], key=lambda n: (n['date'] or '', str(n['note_no']))):
    if n.get('usable') is False: continue
    on = next((g['onSheet'] for g in INV['green'] if str(g['no']) == str(n['note_no']) and g['date'] == n['date']), None)
    put(s, 0, [n['note_no'], n['date'], n['labour_hours'], n.get('cost_rate'), 'yes' if on else ('event week: not billed yet' if (n['date'] or '') >= '2026-10-19' else 'NO'), None, n['cost']], fmt={3: '$#,##0.00', 6: M})
put(s, 0, ['Total', None, S['green_hours'], None, None, None, d(GREEN)], H, GRN, fmt={6: M})
put(s, 0, [])
put(s, 0, ["Extras on Advanced's purchase orders"], H, BLU)
put(s, 0, ['Order', 'Date', 'What', 'Qty', 'Unit', 'Rate', 'We pay'], H)
for x in sorted(EX, key=lambda x: (x['po'], x['date'])):
    put(s, 0, [x['po'], x['date'], x['name'], x['qty'], x['unit'], x['rate'], d(x['cents'])], fmt={5: '$#,##0.00', 6: M})
put(s, 0, ['Total', None, None, None, None, None, d(TL + MAT)], H, GRN, fmt={6: M})

out = HERE / 'GC500_Fencing_Charge_and_Cost_25Oct2026.xlsx'; wb.save(out)
summary = {'dockets': len(use), 'charged': d(TC), 'paid_dockets': d(TP), 'team_leader': d(TL), 'materials': d(MAT), 'invoice_adjustment': d(ADJ), 'fencing_cost': d(FENCE_COST),
           'labour_notes': d(GREEN), 'paid_all_in': d(ALL_IN), 'other_advanced_invoices': d(OTHER_C), 'paid_everything': d(EVERYTHING), 'margin': d(TC - EVERYTHING),
           'orders': [dict(zip(['week', 'order', 'sheet', 'order_amount', 'receipt', 'match', 'carried'], r)) for r in ORDERS],
           'items': {k: {'charge_qty': charge[k]['qty'], 'charge_rate': charge[k]['rate'], 'charge': d(charge[k]['c']), 'pay_qty': pay[k]['qty'], 'pay_rates': sorted(pay[k]['rates']), 'pay': d(pay[k]['c'])} for k in ITEMS},
           'scrim_once': {'invoiced_dockets': len(check), 'clean_is_paper_less_scrim_and_gates': exact, 'scrim_m': SCRIM_M, 'paper_clean_m': PAPER_CLEAN,
                          'taken_off_the_charge': d(ONCE), 'charge_before': d(TC + ONCE), 'baseplan': d(BASEPLAN), 'baseplan_short': d(TC - BASEPLAN)},
           'by_week': {k: {kk: (d(v) if kk in ('charge', 'pay', 'extras', 'notes', 'adj') else v) for kk, v in w.items()} for k, w in W.items()}}
(HERE / 'evidence' / 'fence_pay_charge.json').write_text(json.dumps(summary, indent=1))
print(json.dumps({k: summary[k] for k in ('dockets', 'charged', 'paid_dockets', 'team_leader', 'materials', 'invoice_adjustment', 'fencing_cost', 'labour_notes', 'paid_all_in', 'other_advanced_invoices', 'paid_everything', 'margin')}))
for r in ORDERS: print(r)
print(json.dumps(summary['scrim_once']))
for k, v in summary['items'].items(): print(k, v)
for k, v in summary['by_week'].items(): print(k, v)
