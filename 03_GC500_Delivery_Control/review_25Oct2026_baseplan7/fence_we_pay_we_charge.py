# Author: Andrew Fisher. 25 Oct 2026 - the fencing: what we pay Advanced, then what we charge the V8s, one quantity per line.
# Andrew: "what we are being invoiced for is what we will charge V8s for simple" and "i dont want a story i want what we pay
# and then what we charge simple" (after "we charge v8 based off what we pay advance", "dont change numbers", "should be same
# qty"). Each line is what Advanced invoice us for (their approved sheets for weeks 1 to 5; week 6 and the event week are our
# dockets at their rates until their breakdown comes); the V8s are charged the same quantity at the 2026 card or Andrew's named
# rates. Reads the read-only dumps of the live page (dump_fence_pay_charge.cjs, dump_fence_split.cjs, dump_fence_invoices.cjs);
# writes GC500_Fencing_We_Pay_We_Charge_25Oct2026.xlsx. Values only; both totals are checked against the page to the cent (v12.19:
# the page charges Advanced's quantities, so the charge column is the page's own lines).
#   python3 fence_we_pay_we_charge.py <fence_pay_charge.json> <fence_split.json> <fence_invoices.json>
import collections, json, sys
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.page import PageMargins

D, SPLIT, INV = (json.load(open(a)) for a in sys.argv[1:4])
HERE = Path(__file__).resolve().parent
cents = lambda v: int(Decimal(str(v or 0)).scaleb(2).quantize(Decimal(1), ROUND_HALF_UP))
d = lambda c: c / 100
use = [x for x in D['dockets'] if x['usable']]
RATE = {c['key']: c['charge']['value'] for c in D['columns']}
q = lambda lines, *cols: sum(l['qty'] for l in lines if l['column'] in cols)
CCB = ('ccb_demarc', 'ccb_event', 'flat_feet')

# --- what Advanced invoice us for, by item: their approved sheet where an order carries one, else our docket at their rates.
# Their CCB is one line at one rate, so its metres take our dockets' types (where their sheet has less, the type they left off).
Q = collections.defaultdict(float); P = collections.defaultdict(int); PR = collections.defaultdict(set)
for x in use:
    ours = x['lines']
    if x['supplier']:
        th = x['supplier']['lines']
        for l in th:
            if l['column'] in CCB: continue
            Q[l['column']] += l['qty']; P[l['column']] += l['cents']; PR[l['column']].add(l['rate'])
        short = q(ours, *CCB) - q(th, *CCB); mine = {k: q(ours, k) for k in CCB}; assert short >= 0, x['no']
        for k in ('flat_feet', 'ccb_event', 'ccb_demarc'): cut = min(mine[k], short); mine[k] -= cut; short -= cut
        assert sum(l['cents'] for l in th if l['column'] in CCB) == sum(cents(mine[k] * 5.4) for k in CCB), x['no']
        for k in CCB: Q[k] += mine[k]; P[k] += cents(mine[k] * 5.4); PR[k].add(5.4) if mine[k] else None
    else:
        for l in ours:
            Q[l['column']] += l['qty']
            if l.get('paid') is not None: P[l['column']] += cents(l['paid']); PR[l['column']].add(l['paid_rate'])
assert sum(P.values()) == cents(sum(x['paid_total'] for x in use)), (sum(P.values()), sum(x['paid_total'] for x in use))
S, O, EX = SPLIT['split'], SPLIT['orders'], SPLIT['extras']
tl = [x for x in EX if x['column'] == 'team_leader']; mats = [x for x in EX if x['column'] != 'team_leader']
ADJ, GREEN = cents(O['adjustment']), cents(S['green'])
OUT_IN = cents(S.get('outside_orders1218') or 0)   # v12.18: the orders outside the programme weeks are inside the page's fencing cost
assert sum(P.values()) + sum(x['cents'] for x in EX) + ADJ + OUT_IN == cents(S['paid'])
ADV = [o for o in INV['pos'] if o['po'].get('supplier_name') == 'Advanced Temporary Fencing']
OTHER = [o for o in ADV if not o['weeks']]
PAID = cents(S['paid']) + GREEN + sum(cents(o['po']['amount']) for o in OTHER) - OUT_IN
one = lambda k: (lambda r: (assert_one(k, r), r.pop())[1])(set(PR[k]))
def assert_one(k, r): assert len(r) == 1, (k, r)
assert not Q['fence_blocks'] and Q['relocation'] == 2 and not Q['removal']   # the blocks are not on their invoice; relocation is their hours

# --- the lines: the quantity we pay for, our rate on the same quantity
GATE = float(Decimal(str(2.5 * RATE['clean'])).quantize(Decimal('0.01'), ROUND_HALF_UP))   # a vehicle gate is the 2.5 m of fence it takes the place of
LINES = [('Clean fence', 'clean', 'm', RATE['clean']), ('Scrim fence', 'scrim', 'm', RATE['scrim']), ('CCB demarcation', 'ccb_demarc', 'm', RATE['ccb_demarc']),
         ('CCB event', 'ccb_event', 'm', RATE['ccb_event']), ('CCB flat feet', 'flat_feet', 'm', RATE['flat_feet']), ('Labour', 'labour', 'h', RATE['labour']),
         ('Relocation', 'relocation', 'h', RATE['labour']), ('Pedestrian gates', 'ped_gates', 'each', round(RATE['ped_gates'], 4)), ('Vehicle gates', 'v_gates', 'each', GATE),
         ('Bases', 'bases', 'each', RATE['fence_blocks'])]
# what we charge: the page's own charge lines (v12.19 charges Advanced's quantities), docket by docket, to the cent
row_of = lambda l: ('relocation' if l.get('name') == 'Relocation labour' else 'labour') if l['column'] == 'labour' else ('bases' if l['column'] == 'fence_blocks' and l.get('name') == 'Bases' else l['column'])
CH = collections.defaultdict(int); CQ = collections.defaultdict(float)
for x in use:
    for l in x['lines']:
        if l.get('cost') is None: continue
        CH[row_of(l)] += cents(l['cost']); CQ[row_of(l)] += l['qty']
ROWS = [[name, Q[k], unit, one(k), P[k], rate, CH[k]] for name, k, unit, rate in LINES]
for r, (name, k, unit, rate) in zip(ROWS, LINES): assert r[4] == cents(r[1] * r[3]) and round(CQ[k], 2) == round(Q[k], 2), (r, CQ[k])
assert set(k for k in CH if CH[k]) <= set(k for _, k, _, _ in LINES), dict(CH)
B = {l['id']: cents(l['amount']) for l in (D.get('billed') or {}).get('lines', [])}
ROWS.append(['Team leader', sum(x['qty'] for x in tl), 'days', 200, sum(x['cents'] for x in tl), 200, B['TEAM-LEADER']])
for x in mats: ROWS.append([x['name'].split(' supply')[0], x['qty'], x['unit'], x['rate'], x['cents'], x['rate'], B[x['name'].split(' supply')[0].upper().replace(' ', '-')]])
ROWS.append(['Labour notes', S['green_hours'], 'h', S['green_rate'], GREEN, RATE['labour'], B['LABOUR-NOTES']])
for o in OTHER: ROWS.append(['Partial collection', 1, 'each', o['po']['amount'], cents(o['po']['amount']), o['po']['amount'], B['ORDER-' + o['po']['number']]])
ROWS.append(['Invoice adjustment, weeks 5 and 6', None, None, None, ADJ, None, None])
assert B['TEAM-LEADER'] == ROWS[len(LINES)][4] and B['LABOUR-NOTES'] == cents(S['green_hours'] * RATE['labour']), B
PAY, CHARGE = sum(r[4] for r in ROWS), sum(r[6] or 0 for r in ROWS)
assert PAY == PAID, (PAY, PAID)
assert CHARGE == cents(D['derived']['charged']), (CHARGE, D['derived']['charged'])   # the Fencing tab's figure, to the cent

# --- by invoice (what we pay only)
WEEK_OF_SHEET = {'CON WK6': 'Week 6', 'CON WK5': 'Week 5', 'CON WK4': 'Week 4', 'CON WK3': 'Week 3', 'CON WK2': 'Week 2', 'CON WK1': 'Week 1', 'EVENT WEEK': 'Event Week'}
po_week = {o['po']['number']: WEEK_OF_SHEET.get(o['po'].get('programme_sheet')) for o in ADV}
SPANS = [('Week 6', '2026-09-07', '2026-09-13'), ('Week 5', '2026-09-14', '2026-09-20'), ('Week 4', '2026-09-21', '2026-09-27'), ('Week 3', '2026-09-28', '2026-10-04'),
         ('Week 2', '2026-10-05', '2026-10-11'), ('Week 1', '2026-10-12', '2026-10-18'), ('Event Week', '2026-10-19', '2026-10-25')]
week_of_day = lambda day: next((w for w, a, b in SPANS if a <= (day or '') <= b), None)
W = collections.OrderedDict((w, 0) for w, a, b in SPANS)
for x in use: W[x['week']] += cents(x['paid_total'])
for x in EX: W[po_week[x['po']]] += x['cents']
notes = [n for n in D['notes'] if n.get('usable') is not False]
for n in notes: W[week_of_day(n['date'])] += cents(n['cost'])
for o in O['orders']: W[po_week[o['number']]] += cents(o['adjustment'])
def status(o):
    if o['po'].get('invoice_no'): return 'Invoice received'
    if o['sheet'] and o['sheet'].get('valid'): return 'Approved sheet, no invoice yet'
    return 'Nothing from Advanced yet (our estimate)'
INVOICES = []
for w in W:
    o = next((x for x in ADV if po_week.get(x['po']['number']) == w), None)
    if o: INVOICES.append([status(o), w, o['po']['number'], o['po'].get('amount'), o['po'].get('invoice_no'), d(W[w])])
for o in OTHER: INVOICES.append([status(o), 'Partial collection', o['po']['number'], o['po'].get('amount'), o['po'].get('invoice_no'), o['po'].get('amount')])
ORDER = ['Invoice received', 'Approved sheet, no invoice yet', 'Nothing from Advanced yet (our estimate)']
INVOICES.sort(key=lambda r: (ORDER.index(r[0]), [s[0] for s in SPANS].index(r[1]) if r[1] in [s[0] for s in SPANS] else 9))
BY_STATUS = collections.OrderedDict((s, sum(cents(r[5]) for r in INVOICES if r[0] == s)) for s in ORDER)
assert sum(BY_STATUS.values()) == PAID, (BY_STATUS, PAID)

# ------------------------------------------------------------------------------------------------- workbook
wb = Workbook(); wb.remove(wb.active)
H = Font(bold=True); BIG = Font(bold=True, size=14)
GRN = PatternFill('solid', fgColor='E2EFDA'); BLU = PatternFill('solid', fgColor='DDEBF7'); GRY = PatternFill('solid', fgColor='F2F2F2'); AMB = PatternFill('solid', fgColor='FFF2CC')
M, QF, R4 = '$#,##0.00', '#,##0.0#', '$#,##0.00##'
def sheet(name, widths):
    s = wb.create_sheet(name)
    for i, w in enumerate(widths): s.column_dimensions[chr(65 + i)].width = w
    s.page_setup.orientation = 'landscape'; s.page_setup.fitToWidth = 1; s.page_setup.fitToHeight = 0
    s.sheet_properties.pageSetUpPr.fitToPage = True; s.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5)
    return s
def put(s, vals, font=None, fill=None, fmt=None):
    s.append(list(vals)); r = s.max_row
    for j, c in enumerate(s[r]):
        c.alignment = Alignment(wrap_text=True, vertical='top')
        if font: c.font = font
        if fill and j < len(vals): c.fill = fill
        if fmt and j in fmt: c.number_format = fmt[j]
    return r

s = sheet('Pay and charge', [32, 11, 7, 14, 16, 14, 16])
put(s, ['GC500 fencing: what we pay Advanced and what we charge the V8s'], BIG); put(s, ['AUD ex GST · 25 Oct 2026 · the same quantity both sides; the totals are the page\'s own, to the cent']); put(s, [])
put(s, ['Item', 'Qty', 'Unit', 'We pay (rate)', 'We pay', 'We charge (rate)', 'We charge'], H, BLU)
for name, qty, unit, pr, pc, cr, cc in ROWS:
    put(s, [name, qty, unit, pr, d(pc), cr, None if cc is None else d(cc)], fmt={1: QF, 3: R4, 4: M, 5: R4, 6: M})
put(s, ['Total', None, None, None, d(PAY), None, d(CHARGE)], H, GRN, {4: M, 6: M})
put(s, ['Margin', None, None, None, None, '%.1f%%' % (100 * (CHARGE - PAY) / CHARGE), d(CHARGE - PAY)], H, GRN, {6: M})
s.freeze_panes = 'A5'

s = sheet('Invoices', [36, 18, 11, 14, 11, 16])
put(s, ['What we pay Advanced, by invoice'], BIG); put(s, ['AUD ex GST · 25 Oct 2026']); put(s, [])
put(s, ['Status', 'Week', 'Order', 'Order amount', 'Receipt', 'We pay'], H, BLU)
for st in ORDER:
    for r in [x for x in INVOICES if x[0] == st]: put(s, r, fmt={3: M, 5: M})
    put(s, [st + ', total', None, None, None, None, d(BY_STATUS[st])], H, AMB if 'estimate' in st else GRY, {5: M})
put(s, ['Total we pay Advanced', None, None, None, None, d(PAID)], H, GRN, {5: M})
s.freeze_panes = 'A5'

out = HERE / 'GC500_Fencing_We_Pay_We_Charge_25Oct2026.xlsx'; wb.save(out)
print(json.dumps({'pay': d(PAY), 'charge': d(CHARGE), 'margin': d(CHARGE - PAY), 'by_status': {k: d(v) for k, v in BY_STATUS.items()}, 'file': out.name,
                  'rows': [[r[0], r[1], r[3], d(r[4]), r[5], None if r[6] is None else d(r[6])] for r in ROWS]}))
