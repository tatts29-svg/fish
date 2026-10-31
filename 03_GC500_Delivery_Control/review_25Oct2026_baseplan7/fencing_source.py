# Author: Andrew Fisher. 25 Oct 2026 - where Baseplan's fencing numbers (9990776) come from. Andrew: "where is baseplan
# getting numbers from have not seen the new dockets is that why". Compares the 9990776 lines in each export with the
# signed fencing dockets (split at 15 Oct, when the contract first appeared), the 2026 fencing programme and the export's
# own Delivery Dockets sheet. Reads private files only; writes evidence/fencing_source.json (no names, no links).
#   python3 fencing_source.py <private dir> [<Baseplan (5) xlsx> [<record snapshot json> ...]]
# record snapshots (private): what our record held when 9990776 was keyed, to test "did they copy our total"
# private dir: bp7.json, bp6.json (to_json.py), page_DATA.json (the live page's DATA), fence_detail.json
# (dump_fence_detail.cjs), fence.json (dump_fence.cjs)
import collections, json, sys
from pathlib import Path

PRIV = Path(sys.argv[1]); HERE = Path(__file__).resolve().parent
J = lambda n: json.load(open(PRIV / n))
C = '9990776'
line = lambda l: {'line': l['Line'], 'description': l['Description'], 'quantity': l['Quantity'], 'rate': l.get('Rate 1', l.get('Price')), 'status': l['Status']}

bp7, bp6, D = J('bp7.json'), J('bp6.json'), J('page_DATA.json')
exports = {'(4) 15 Oct': [{'line': r.get('line'), 'description': r.get('description'), 'quantity': r.get('quantity'), 'rate': r.get('rate_1'), 'status': r.get('status_as_written')}
                          for r in D['rental_on_hire']['source_history949']['rows'] if str(r.get('rental_contract')) == C]}
if len(sys.argv) > 2:   # Baseplan (5), 16 Oct, read straight from the workbook
    from openpyxl import load_workbook
    ws = load_workbook(sys.argv[2], read_only=True, data_only=True)[C + '-STPS']
    rows = list(ws.iter_rows(values_only=True)); h = next(i for i, r in enumerate(rows) if r and 'Description' in r and 'Quantity' in r); H = rows[h]
    exports['(5) 16 Oct'] = [{'line': r[H.index('Line')], 'description': r[H.index('Description')], 'quantity': r[H.index('Quantity')], 'rate': r[H.index('Rate 1')] if 'Rate 1' in H else None,
                              'status': r[H.index('Status')]} for r in rows[h + 1:] if r and r[H.index('Description')]]
exports['(6) 22 Oct'] = [line(l) for l in bp6[C]]
exports['(7) 25 Oct'] = [line(l) for l in bp7[C]]
fence_q = lambda L: {x['line']: x['quantity'] for x in L if x['line'] in range(2, 10)}   # lines 2-9; line 1 is WB13, line 10 shade cloth
same = all(fence_q(L) == fence_q(exports['(7) 25 Oct']) for L in exports.values())

# the signed dockets the P&L charges, split at 15 Oct
R = {k: (v or {}).get('value') or 0 for k, v in J('fence.json')['rates'].items()}
split = {}
for x in J('fence_detail.json')['dockets']:
    if not x['usable']: continue
    h = 'worked to 15 Oct' if (x['date'] or '0000') <= '2026-10-15' else 'worked 16-23 Oct'
    s = split.setdefault(h, {'dockets': 0, 'quantities': collections.defaultdict(float), 'at_card': 0.0})
    s['dockets'] += 1
    for k, v in (x['quantities'] or {}).items():
        s['quantities'][k] += float(v or 0); s['at_card'] += float(v or 0) * R.get(k, 0)
for s in split.values(): s['quantities'] = {k: round(v, 2) for k, v in sorted(s['quantities'].items())}; s['at_card'] = round(s['at_card'], 2)
ever = collections.defaultdict(float)
for s in split.values():
    for k, v in s['quantities'].items(): ever[k] += v

# the 2026 fencing programme (received 9 Oct): build and event weeks, as the workbook has them
prog = collections.defaultdict(float)
for w in D['fencing']['week_sheets']:
    if w['phase'] != 'Deconstruction':
        for k, v in (w.get('programme_totals') or {}).items(): prog[k] += v or 0

# our own record when the contract was keyed: the signed fencing dockets in by then (service notes are not red-book charges)
then = []
for snap in sys.argv[3:]:
    S = json.load(open(snap)); t = collections.defaultdict(float); n = 0
    for v in (S.get('docs', {}).get('fenceDockets') or {}).values():
        if not isinstance(v, dict) or v.get('record_type') == 'service-work': continue
        n += 1
        for k, q in (v.get('quantities') or {}).items(): t[k] += float(q or 0)
    then.append({'record': S.get('version'), 'dockets': n, 'quantities': {k: round(q, 2) for k, q in sorted(t.items())}})

# the export's own Delivery Dockets sheet: Coates' trucks, not Advanced's fencing dockets
dd = [r for r in bp7['Delivery Dockets'] if str(r.get('Hire Agreement')) == C]
dates = sorted(str(r.get('Date Delivery'))[:10] for r in bp7['Delivery Dockets'] if r.get('Date Delivery'))

bp_total = round(sum((l['Quantity'] or 0) * (l.get('Rate 1') or l.get('Price') or 0) for l in bp7[C] if l['Line'] != 1), 2)
out = {'contract': C, 'baseplan_lines_unchanged_since_15_Oct': same, 'exports': exports,
       'baseplan_fencing_total': bp_total, 'dockets': split, 'dockets_to_date': {k: round(v, 2) for k, v in sorted(ever.items())},
       'more_on_baseplan_than_every_docket': {k: [b, ever.get(d, 0)] for k, b, d in (('crowd control barrier m', fence_q(exports['(7) 25 Oct'])[7], 'ccb_event'), ('ped gates', fence_q(exports['(7) 25 Oct'])[2], 'ped_gates')) if b > ever.get(d, 0)},
       'programme_2026_build_and_event': {k: round(v, 2) for k, v in sorted(prog.items())}, 'our_record_then': then,
       'export_delivery_dockets': {'rows': len(bp7['Delivery Dockets']), 'first': dates[0], 'last': dates[-1], 'same_as_22_Oct': bp7['Delivery Dockets'] == bp6.get('Delivery Dockets'),
                                   'on_9990776': [{'docket': r.get('Delivery Docket'), 'date': str(r.get('Date Delivery'))[:16], 'direction': r.get('Direction')} for r in dd]}}
(HERE / 'evidence' / 'fencing_source.json').write_text(json.dumps(out, indent=1, default=str))
print(json.dumps({'unchanged': same, 'exports': list(exports), 'baseplan_total': bp_total, 'split': {h: [s['dockets'], s['at_card']] for h, s in split.items()},
                  'more_on_baseplan': out['more_on_baseplan_than_every_docket'], 'delivery_dockets': {k: v for k, v in out['export_delivery_dockets'].items()}}, default=str))
