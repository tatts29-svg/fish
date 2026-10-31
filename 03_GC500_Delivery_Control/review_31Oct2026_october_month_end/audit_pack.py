"""Author: Andrew Fisher. Independent audit of the GC500 P&L pack: recomputes the figures from the source data (not from the
builder) and checks the workbook, page and CSV agree with them and with each other, to the cent. Read only.

    python3 -I audit_pack.py <pack_dir> <progress_dir> [<alfie_nights.json>]

<pack_dir> holds GC500_PL_final.xlsx/.html, GC500_October_month_end.csv and GC500_labour_by_month_and_branch.csv;
<progress_dir> holds the private inputs (invoices.json, dated_costs.json, cost_sources.json, bp12.json, dated_all.json).
"""
import csv, html, json, re, sys, collections
from pathlib import Path
import openpyxl

PK, PR = Path(sys.argv[1]), Path(sys.argv[2])
ALF = Path(sys.argv[3]) if len(sys.argv) > 3 else None
r2 = lambda x: round(x + 1e-9, 2)
fmt = lambda x: '{:,.2f}'.format(x)
OK, BAD = [], []
def check(name, cond, detail=''):
    (OK if cond else BAD).append(name + ('' if cond else '  <-- ' + str(detail)))
def eq(name, a, b, tol=0.005):
    check(name, abs((a or 0) - (b or 0)) < tol, '%s vs %s' % (a, b))

wb = openpyxl.load_workbook(PK / 'GC500_PL_final.xlsx', data_only=True)
H = open(PK / 'GC500_PL_final.html', encoding='utf-8').read()
HT = re.sub(r'\s+', ' ', html.unescape(re.sub(r'<[^>]+>', ' ', H)))
CS = list(csv.reader(open(PK / 'GC500_October_month_end.csv', encoding='utf-8-sig')))
LAB = list(csv.reader(open(PK / 'GC500_labour_by_month_and_branch.csv', encoding='utf-8-sig')))
rows = lambda name: [r for r in wb[name].iter_rows(min_row=2, values_only=True)]

# ------------------------------------------------------------------ revenue, from the invoices
inv = json.load(open(PR / 'invoices.json')); inv = inv if isinstance(inv, list) else list(inv.values())
check('19 invoices in the source', len(inv) == 19, len(inv))
eq('the 19 invoices total', sum(i['ex_gst'] for i in inv), 1003037.12)
OUT = {'25124396'}                                          # invoiced to Gold Coast Helitours, not V8 Supercars
FIX = {'25120652': 'STPS', '25120689': 'STPS'}              # invoice header "Branch Staplyton (S)", Sales Analysis Code STPS-CON
keep = [i for i in inv if i['invoice'] not in OUT]
REV_TOT = r2(sum(i['ex_gst'] for i in keep)); eq('18 invoices to V8 Supercars', REV_TOT, 1001063.07)
REVB = collections.defaultdict(float)
for i in keep: REVB[FIX.get(i['invoice'], i.get('bp_branch'))] += i['ex_gst']
rv = [r for r in rows('Revenue by invoice') if r and r[0] and r[3] != 'removed']
check('revenue sheet: 18 invoices', len(rv) == 18, len(rv))
src = {i['invoice']: i for i in keep}
for r in rv:
    i = src.get(r[0]); check('revenue sheet row is a source invoice: %s' % r[0], i is not None)
    if not i: continue
    eq('ex GST matches the invoice: %s' % r[0], r[4], i['ex_gst'])
    eq('months + services = ex GST: %s' % r[0], sum((r[k] or 0) for k in (6, 7, 8, 9)), r[4])
    check('branch: %s' % r[0], r[3].split(' ')[0] == FIX.get(i['invoice'], i.get('bp_branch')), r[3])
rem = [r for r in rows('Revenue by invoice') if r and r[3] == 'removed']
check('25124396 shown as removed, with no amount', len(rem) == 1 and rem[0][0] == '25124396' and all(x is None for x in rem[0][4:10]), rem)
summ = {r[0]: r for r in rows('Summary') if r and r[0] in ('KINP', 'STPS', 'NVAC', 'MEAD', 'The job')}
for b in ('KINP', 'STPS', 'NVAC', 'MEAD'): eq('branch revenue %s' % b, summ[b][1], REVB[b])
eq('revenue: invoices + November forecast', summ['The job'][3], REV_TOT + 839.07)
check('the 839.07 forecast sits on NVAC', abs(summ['NVAC'][2] - 839.07) < 0.005, summ['NVAC'][2])

# ------------------------------------------------------------------ costs, from the record's dated pieces
cs = [r for r in rows('Costs by supplier') if r and r[0]]
col = lambda i: r2(sum((r[i] or 0) for r in cs))
SEP, OCT, NOV, ACC = col(7), col(8), col(9), col(10)
JOB = r2(SEP + OCT + NOV)
dc = json.load(open(PR / 'dated_costs.json'))
T = collections.defaultdict(float)
for t in dc['transport']:
    mo = (t['d'] or '')[:7]; T[('rec', mo)] += t.get('actual') or 0; T[('plan' if t.get('planned') else 'fc', mo)] += t.get('f') or 0
car = [r for r in cs if r[0].startswith('Carriers')]
want = [(T[('rec', '2026-09')], T[('rec', '2026-10')], T[('rec', '2026-11')]), (T[('fc', '2026-09')], T[('fc', '2026-10')], T[('fc', '2026-11')]), (T[('plan', '2026-09')], T[('plan', '2026-10')], T[('plan', '2026-11')])]
check('three carrier rows', len(car) == 3, len(car))
for k, (r, w) in enumerate(zip(car, want)):
    for j, mo in enumerate(('Sep', 'Oct', 'Nov')): eq('carriers row %d %s = the transport rows by date' % (k + 1, mo), r[7 + j], w[j])
acc = [x for x in dc['exp'] if x['k'] == 'accommodation' and x['amount']]
A = {m: r2(sum(x['amount'] for x in acc if x['d'][:7] == m)) for m in ('2026-09', '2026-10', '2026-11')}
eq('tracker accommodation Sep', A['2026-09'], 6370.36); eq('tracker accommodation Oct', A['2026-10'], 11851.89); eq('tracker accommodation Nov', A['2026-11'], 3801.71)
accr = [r for r in cs if r[0].startswith('Accommodation')]
k97 = [r for r in accr if r[3].startswith('97 priced nights')][0]; lud = [r for r in accr if r[3].startswith('Ludwig')][0]; alf = [r for r in accr if r[3].startswith('Alfie')][0]
eq('97 nights row Oct = tracker Oct less Ludwig 655.71', k97[8], A['2026-10'] - 655.71); eq('Ludwig row = 3 x 218.57', lud[8], 3 * 218.57)
check('Ludwig row on NVAC', lud[1] == 'NVAC', lud[1])
if ALF:
    al = json.load(open(ALF))['alfie']
    un = [x for x in al if x['amt'] is None]
    late = [x['amt'] for x in al if x['amt'] is not None and ('2026-10-26' <= x['d'] <= '2026-10-29' or '2026-11-02' <= x['d'] <= '2026-11-05')]
    check("Alfie's rate on 26-29 Oct and 2-5 Nov is 223.63 on all 8 nights", len(late) == 8 and set(late) == {223.63}, late)
    n_oct = sum(1 for x in un if x['d'][:7] == '2026-10'); n_nov = sum(1 for x in un if x['d'][:7] == '2026-11')
    check('Alfie unpriced nights: 8 Oct, 4 Nov', (n_oct, n_nov) == (8, 4), (n_oct, n_nov))
    eq('Alfie Oct = 8 x 223.63', alf[8], r2(n_oct * 223.63)); eq('Alfie Nov = 4 x 223.63', alf[9], r2(n_nov * 223.63)); eq('Alfie accrues October only', alf[10], r2(n_oct * 223.63))
meals = [x for x in dc['exp'] if x['k'] in ('meals', 'misc') and x['amount']]
M = {m: r2(sum(x['amount'] for x in meals if x['d'][:7] == m)) for m in ('2026-09', '2026-10', '2026-11')}
mr = [r for r in cs if r[0].startswith('Meals')]
eq('meals Sep', sum((r[7] or 0) for r in mr), M['2026-09']); eq('meals Oct', sum((r[8] or 0) for r in mr), M['2026-10']); eq('meals Nov', sum((r[9] or 0) for r in mr), M['2026-11'])
SALH = collections.defaultdict(float)
for r in LAB[2:]:
    if r and r[0][:4] == '2026' and r[3] == 'Salary' and '2026-09-14' <= r[0] <= '2026-11-13': SALH[r[0][5:7]] += float(r[4] or 0)
tot_h = sum(SALH.values()); sal = [r for r in cs if r[0].startswith('Salary')][0]
eq('salary Sep (in-window hours)', sal[7], r2(14700 * SALH['09'] / tot_h)); eq('salary Oct', sal[8], r2(14700 * SALH['10'] / tot_h)); eq('salary total 14,700.00', sal[7] + sal[8] + sal[9], 14700)
LHW = collections.defaultdict(float)
import datetime as dt
for r in LAB[2:]:
    if r and r[0][:4] == '2026' and r[3] == 'Labour hire' and r[5] not in ('', 'hours only'):
        d = dt.date.fromisoformat(r[0]); we = d + dt.timedelta(days=(6 - d.weekday())); LHW[we.isoformat()] += float(r[5])
jc = {r[3].replace('Running sheet ', ''): r for r in cs if r[0].startswith('Job Connect')}
for we, key in (('2026-10-11', 'w/e 11 Oct'), ('2026-10-18', 'w/e 18 Oct'), ('2026-10-25', 'w/e 25 Oct'), ('2026-11-01', 'w/e 1 Nov'), ('2026-11-08', 'w/e 8 Nov')):
    eq('Job Connect %s = running sheet week' % key, sum((jc[key][k] or 0) for k in (7, 8, 9)), LHW[we])
pos = {o['number']: o for o in json.load(open(PR / 'cost_sources.json'))['pos']}
for po, key in (('4657010', 'w/e 20 Sep'), ('4660565', 'w/e 27 Sep'), ('4661783', 'w/e 4 Oct')):
    eq('Job Connect %s = PO %s' % (key, po), sum((jc[key][k] or 0) for k in (7, 8, 9)), float(pos[po]['amount']))
eq('Advanced invoiced POs = 214,225.00', sum(float(pos[p]['amount']) for p in ('4647508', '4647509', '4647515', '4647517', '4647519', '4647520', '4658850')), 214225.00)
adv = [r for r in cs if r[0].startswith('Advanced')]
eq('Advanced invoiced row = the POs', sum((adv[0][k] or 0) for k in (7, 8, 9)), 214225.00)
eq('Event Portables = the four quotes', sum((r[7] or 0) + (r[8] or 0) + (r[9] or 0) for r in cs if r[0] == 'Event Portables'), 42175.00 + 51372.00 + 4850.00 + 3350.00)
for po in ('4655573', '4663216', '4666720', '4670270', '4670876', '4669481'):
    row = [r for r in cs if r[4] == po][0]; eq('consumables PO %s at its amount' % po, (row[7] or 0) + (row[8] or 0) + (row[9] or 0), float(pos[po]['amount']))
check('every PO on the record is in the pack', all(any(p in str(r[4]) for r in cs) for p in pos), [p for p in pos if not any(p in str(r[4]) for r in cs)])

# ------------------------------------------------------------------ the pack's own cross-foots
cb = [r for r in rows('Costs by branch') if r and r[0]]
CBB = collections.defaultdict(float)
for r in cb: CBB[r[0]] += r[2]
eq('costs by supplier = costs by branch', JOB, sum(CBB.values()))
for b in ('KINP', 'STPS', 'NVAC', 'MEAD'): eq('summary costs %s = costs by branch' % b, summ[b][4], CBB[b])
eq('summary job cost = costs by supplier', summ['The job'][4], JOB)
eq('difference = revenue - costs', summ['The job'][5], summ['The job'][3] - JOB)
mu = [r for r in rows('Match-up by contract') if r]
job_row = [r for r in mu if r[0] == 'The job'][0]
eq('match-up: tied + untied = job cost', job_row[7], JOB)
eq('match-up: the costs listed add to the job cost', sum((r[7] or 0) for r in mu if r[0] != 'The job'), JOB)
eq('match-up: invoices = 18-invoice revenue', job_row[3], REV_TOT)
check('match-up: 18 contracts', len({r[0] for r in mu if r[0] and r[0][0].isdigit()}) == 18)
jn = [r for r in rows('Journals') if r and r[0]]
accr_j = [r for r in jn if str(r[0]).startswith('Accrual')]
eq('journals = the accrue column', sum(r[5] for r in accr_j), ACC)
check('journals total row', any(r[0] == 'Total accruals' and abs(r[5] - ACC) < 0.005 for r in jn))
moves = [r for r in jn if str(r[0]).startswith('Branch move ·')]   # the two moves, not the header row above them
check('branch moves net to zero (two legs each)', all(re.search(r'−([\d,.]+) · \w+ \+([\d,.]+)', r[4]) and re.search(r'−([\d,.]+)', r[4]).group(1) == re.search(r'\+([\d,.]+)', r[4]).group(1) for r in moves) and len(moves) == 2, [r[4] for r in moves])
lp = [r for r in rows('Labour by person') if r and r[0]]
eq('labour hours = the labour register (2,304.78)', sum((r[6] or 0) + (r[7] or 0) + (r[8] or 0) for r in lp), sum(float(r[4] or 0) for r in LAB[2:] if r and r[0][:4] == '2026'))
eq('labour hire $ by person = running sheet', sum((r[9] or 0) for r in lp), sum(float(r[5]) for r in LAB[2:] if r and r[0][:4] == '2026' and r[3] == 'Labour hire' and r[5] not in ('', 'hours only')))
eq('P&L labour hire (invoices where received) = 34,748.64', sum((r[7] or 0) + (r[8] or 0) + (r[9] or 0) for r in cs if r[0].startswith('Job Connect')), 34748.64)

# ------------------------------------------------------------------ the CSV and the page say the same
tot_csv = [r for r in CS if r and r[0] == 'Total'][0]
eq('CSV Sep total', float(tot_csv[7]), SEP); eq('CSV Oct total', float(tot_csv[8]), OCT); eq('CSV Nov total', float(tot_csv[9]), NOV); eq('CSV accrue total', float(tot_csv[10]), ACC)
for label, v in (('job cost', JOB), ('accruals', ACC), ('revenue', summ['The job'][3]), ('18 invoices', REV_TOT), ('difference', summ['The job'][5]), ('Sep costs', SEP), ('Oct costs', OCT), ('Nov costs', NOV)):
    check('page shows %s %s' % (label, fmt(v)), fmt(v) in HT, fmt(v))
check('no unformatted placeholders on the page', not re.search(r'\{[a-zA-Z_]+\(', H))
prov = re.compile(r"your word|your call|your rule|\byou\b|\byour\b|ruling|on Andrew|Andrew’s word|Andrew's word|\(Andrew,|Codex|Claude|19 invoices", re.I)
texts = [HT] + [c for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row if isinstance(c, str)] + [c for row in CS for c in row]
bad = sorted({m.group(0) for t in texts for m in prov.finditer(t)})
check('facts only: no "your word", "your call", "Andrew, <date>", agent names or "19 invoices" in the outputs', not bad, bad)

print('AUDIT: %d checks, %d pass, %d fail' % (len(OK) + len(BAD), len(OK), len(BAD)))
print('job cost %s · Sep %s · Oct %s · Nov %s · accruals %s · revenue %s · difference %s' % (fmt(JOB), fmt(SEP), fmt(OCT), fmt(NOV), fmt(ACC), fmt(summ['The job'][3]), fmt(summ['The job'][5])))
for b in BAD: print('FAIL', b)
sys.exit(1 if BAD else 0)
