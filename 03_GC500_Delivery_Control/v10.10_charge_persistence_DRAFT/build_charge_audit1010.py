"""Author: Andrew Fisher. Create a readable private review from the final read-only snapshot."""
import collections
import datetime
import json
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

import os
P=Path(os.environ.get('AUDIT_DIR','/workspace/private-gc500-finance1010-final'))
s=json.loads((P/'coverage-snapshot.json').read_text())
equivalence=json.loads((P/'release-equivalence.json').read_text()) if (P/'release-equivalence.json').exists() else None
release_hash=equivalence['release'] if equivalence and equivalence['auditedSource']==s['sha256'] and equivalence['scriptsIdentical'] else s['sha256']
items={(i['ref'],i['item']):i for i in s['items']}
accessories={(i['ref'],'Accessory · '+x['type']):x for i in s['items'] if i['received'] for x in i.get('accessories') or []}
wb=Workbook();existing=wb.active;existing.title='Existing charges';fix=wb.create_sheet('Needs fixing')
wb.properties.creator='Andrew Fisher';wb.properties.title='GC500 charge coverage — 14 Oct 2026'
wb.properties.subject='Arrived-equipment customer charges and unresolved coverage'
dark='1A303B';orange='F47821';amber='FFF0CF';grey='52626C';green='DBEFE3'
money='"$"#,##0.00;[Red]("$"#,##0.00);"$"0.00'
when=datetime.datetime.fromisoformat(s['at'].replace('Z','+00:00')).astimezone(datetime.timezone(datetime.timedelta(hours=10))).strftime('%d %b %Y %H:%M AEST')

def start(ws,title,sub,headers,widths):
    n=len(headers);ws.sheet_view.showGridLines=False
    for row,text in [(1,title),(2,sub),(3,f'Author: Andrew Fisher · {when} · Shared record {s["nativeVersion"]}'),(4,'Amounts ex GST. Charge inclusion does not mean physical demob, pickup or cleaning is complete.')]:
        ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=n)
        c=ws.cell(row,1,text);c.font=Font(name='Aptos',size=16 if row==1 else 10,bold=row==1,color='FFFFFF' if row==1 else grey)
        c.fill=PatternFill('solid',fgColor=dark if row==1 else 'F3F6F7');c.alignment=Alignment(vertical='center',wrap_text=True)
        ws.row_dimensions[row].height=28 if row==1 else 27
    for k,(name,width) in enumerate(zip(headers,widths),1):
        c=ws.cell(6,k,name);c.font=Font(name='Aptos',size=10,bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor=dark);c.alignment=Alignment(wrap_text=True,vertical='center')
        ws.column_dimensions[c.column_letter].width=width
    ws.row_dimensions[6].height=30;ws.freeze_panes='C7';ws.sheet_view.zoomScale=85
    ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    ws.print_title_rows='1:6';ws.oddFooter.center.text='Author: Andrew Fisher · &P / &N';ws.oddFooter.right.text='AUD ex GST'
    ws.sheet_properties.outlinePr.summaryRight=False

def finish(ws,table_name,end_col):
    table=Table(displayName=table_name,ref=f'A6:{end_col}{ws.max_row}')
    table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(table)
    for row in ws.iter_rows(min_row=7):
        for c in row:
            c.font=Font(name='Aptos',size=10,color=dark);c.alignment=Alignment(vertical='top',wrap_text=True)
        ws.row_dimensions[row[0].row].height=34
    ws.print_options.horizontalCentered=True;ws.print_area=f'A1:{end_col}{ws.max_row}'

start(existing,'Existing charges — equipment received',
    f'Priced onsite subtotal: ${s["bill"]["known"]:,.2f} · Unresolved items are listed on Needs fixing. This is not an issued invoice.',
    ['Reference','Description','Qty recorded','Branch','Hire','Work charges','Delivery + pickup','Priced subtotal','Status / basis'],[13,28,12,11,16,17,20,19,47])
for r in s['bill']['rows']:
    i=items.get((r['ref'],r['item']),{});pending=r['pending'];hire=r['hire'];work=r['work'];transport=r['transport']
    if any(x.startswith('Hire')for x in pending) and hire==0:hire=None
    if any('work' in x.lower() or 'service' in x.lower()for x in pending) and work==0:work=None
    if any(x in ('Delivery','Pickup')for x in pending) and transport==0:transport=None
    basis='Needs fixing: '+', '.join(pending) if pending else 'Priced / linked'
    if any(x.get('ownUse')for x in (i.get('hire')or{}).get('lines',[])):basis='Coates own use · no customer hire'+(' · '+basis if pending else '')
    elif any(x.get('included')for x in (i.get('hire')or{}).get('lines',[])):basis='Hire included in main item'+(' · '+basis if pending else '')
    existing.append([r['ref'],r['item'],r['arrived'],r['branch'],hire,work,transport,r['total'],basis])
    row=existing.max_row
    for col in range(5,9):existing.cell(row,col).number_format=money
    if i:
        services=['%s: %s × %s = %s'%(x['name'],x['rate'],min(i['received'],x['qty'])if x['qty'] is not None else 'unknown',round(x['rate']*min(i['received'],x['qty']),2)if x['qty']is not None and isinstance(x['rate'],(float,int))else 'unknown')for x in i['charges']['rows']if x['qty']]
        existing.cell(row,6).comment=Comment('\n'.join(services)or'No priced applicable work line. Check Needs fixing for unmatched scope.','Andrew Fisher')
    accessory=accessories.get((r['ref'],r['item']))
    if accessory and accessory.get('qty_stated') is False:
        existing.cell(row,3).comment=Comment('Schedule text: '+str(accessory.get('as_written',''))+'. No explicit quantity in that text; shown quantity is the existing model value, not a new onsite count.','Andrew Fisher')
    if pending:existing.cell(row,9).fill=PatternFill('solid',fgColor=amber)
finish(existing,'ArrivedCharges','I')
for row in existing.iter_rows(min_row=7):
    row[3].alignment=Alignment(vertical='top',wrap_text=True,indent=1)
    row[8].alignment=Alignment(vertical='top',wrap_text=True,indent=1)
    for c in row[4:8]:c.alignment=Alignment(horizontal='right',vertical='top',indent=1)

start(fix,'Needs fixing — rates and allocations',
    'Only arrived-equipment questions below. Some charges already exist and need linking; do not add them twice.',
    ['Reference','Description','Branch','Category','Charge / quantity','What needs fixing','Action'],[13,25,11,22,24,47,64])
groups=collections.OrderedDict()
for t in s['transport']['rows']:
    if t['state']!='held' or not t['receivedQuantity1006']:continue
    key=(t['ref'],t['item'],tuple(t['reasons']),t['rate'])
    groups.setdefault(key,[]).append(t)
for (ref,item,reasons,rate),rs in groups.items():
    i=items.get((ref,item),{});legs=' + '.join('Delivery'if r['leg']=='delivery'else'Pickup'for r in rs)
    action='Link the existing contract transport to these movements before adding another charge.'
    if rate is None:action='Confirm the customer rate and whether this is a separate trip or covered by another item.'
    if ref=='GN20':action='Enter the approved delivery and pickup rates in the GN20 reference. Current card says POA.'
    if item=='TL2':action='Confirm customer truck/load rate and billable movements. Do not multiply a bulk-load rate by each barrier.'
    if item=='Waste tank':action='Confirm tank size and whether it travels with the block or separately; link any existing transport.'
    fix.append([ref,item,i.get('branch',{}).get('code'),'Transport Revenue',f'{legs} · {rs[0]["receivedQuantity1006"]} qty','; '.join(reasons),action])
for h in s['bill']['holds']:
    if h['charge']in ('Delivery','Pickup') or h['ref']=='Job':continue
    i=items.get((h['ref'],h['item']),{});branch=i.get('branch',{}).get('code')or next((r['branch']for r in s['bill']['rows']if r['ref']==h['ref']and r['item']==h['item']),'')
    category='Hire allocation';why=h['charge'];action='Match current assets and quantities to existing contract hire. Do not charge the same asset twice.'
    if h['item'].startswith('Accessory'):
        why='Exact type / quantity or contract inclusion not established';action='Confirm the item type and quantity, then link included hire or its applicable rate.'
    if h['ref']=='WC07':
        why='20 received; site number 1317743 differs from contract 9968955 line 38: 1317643'
        action='Confirm the physical number and correct the contract/site link. The contract already carries $90.07; do not add a second charge.'
    if h['ref']=='WC21':
        why='6 received; current 1211977 is not linked. Latest Baseplan 9968955 line 7 has 1211977; app still has 1211963 at WC12.'
        action='Reconcile line 7 and WC12 together using the confirmed site numbers. Keep its existing $90.07 hire once.'
    if h['ref']=='T0001':
        why='Quantity says 8; 7 current numbers. Eighth board 1211404 has returned / rehire history under VMS09.'
        action='Reconcile the old T0001 quantity with VMS09 and contract 9961265 lines 1 / 12. Do not add an eighth hire charge.'
    if h['ref']=='FL02':
        why='1197839 also appears at T0004. Contract 9961976 line 21 already counts its hire there.'
        action='Resolve FL02 / T0004 reference alias. Keep one customer hire charge for asset 1197839.'
    if h['ref']=='FL01':
        why='Asset 50004 is named in contract 9961976 line 31 memo, currently linked to T0003.'
        action='Resolve FL01 / T0003 and contract line 31 linkage before adding any extra hire.'
    if h['ref']=='CP1' and 'Hydration' in h['item']:
        why='Hydration station has a published weekly rate, but this accessory has no linked hire basis.'
        action='Link the station and hire period to Street Rate Card 2026 C50 ($542.8203/week) or its existing contract; avoid duplicate package hire.'
    if h['ref']=='T0021' and 'Swivel' in h['item']:
        why='3 swivel chairs recorded; the card names Office Chair, but the exact product / inclusion is not linked.'
        action='Confirm Office Chair or included building contents, then link the card/contract once.'
    if h['ref']=='T0021' and 'Whiteboard' in h['item']:
        why='No whiteboard rate on the supplied card and no linked contract rate.'
        action='Confirm included contents or an approved separate charge. Do not invent a rate.'
    if 'work'in h['charge'].lower()or'service'in h['charge'].lower():
        category='Work rate / scope';why='No exact service-rate match; empty work list does not mean free';action='Confirm applicable install / demob / other work charge, or explicit inclusion. Do not copy an unrelated rate.'
    fix.append([h['ref'],h['item'],branch,category,h['charge'],why,action])
# A brief separate Direct-cost summary avoids confusing employee cost gaps with customer Revenue.
lab=s['labour'];missing=lab['missingHours'];unpriced=[r for r in s['labourRows']if r['calculatedCost']is None]
fix.append(['LABOUR','Coates / labour hire','Various','Direct costs — labour','Hours missing',f'{len(missing)} shifts: {sum(r["date"]<=s["day"]for r in missing)} past/current; {sum(r["date"]>s["day"]for r in missing)} future','Enter evidenced actual or planned hours. These are Coates costs, separate from customer per-piece charges.'])
fix.append(['LABOUR','Coates / labour hire','Various','Direct costs — labour','Cost rates missing',f'{len(unpriced)} shifts have hours but no usable cost rate','Add approved wage/labour-hire cost rates and review stale approvals; no hours or rates have been invented.'])
finish(fix,'ChargeQuestions','G')
for row in fix.iter_rows(min_row=7):
    fix.row_dimensions[row[0].row].height=55
    row[5].alignment=Alignment(vertical='top',wrap_text=True,indent=1)
    row[6].alignment=Alignment(vertical='top',wrap_text=True,indent=1)
    if row[3].value.startswith('Direct costs'):row[3].fill=PatternFill('solid',fgColor='DCE8F1')
for ws in (existing,fix):ws['A3'].comment=Comment('Release build SHA256: '+release_hash+'\nFinancial capture SHA256: '+s['sha256']+'\nShared record: '+str(s['nativeVersion'])+'\nRead-only coverage audit; no record writes. Source-backed v10.10 audit. Original source documents and unresolved inputs are preserved.','Andrew Fisher')
wb.active=1
out=Path('/workspace/output/GC500_Charge_Audit_14Oct2026.xlsx');out.parent.mkdir(parents=True,exist_ok=True);wb.save(out)
test=load_workbook(out,data_only=False)
assert test.sheetnames==['Existing charges','Needs fixing']
assert existing.max_row-6==len(s['bill']['rows'])
assert round(sum(r[7].value or 0 for r in test['Existing charges'].iter_rows(min_row=7)),2)==s['bill']['known']
assert not any(c.data_type=='f'for ws in test for row in ws for c in row)
summary=dict(author='Andrew Fisher',source=release_hash,financialCapture=s['sha256'],nativeVersion=s['nativeVersion'],output=str(out),existingRows=existing.max_row-6,questionRows=fix.max_row-6,subtotal=s['bill']['known'],noFormulas=True,privateSnapshot=str(P/'coverage-snapshot.json'))
(P/'workbook-validation.json').write_text(json.dumps(summary,indent=2)+'\n')
print(json.dumps(summary))
