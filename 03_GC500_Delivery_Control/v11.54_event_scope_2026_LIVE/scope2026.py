"""Author: Andrew Fisher. v11.54 - the Coates Event Labour Scope C004-26003-01 (23 Oct 2026) transcribed into the page's
event-scope schema (gc500-event-labour-scope-1), checked against the workbook's own cached figures to the cent.

    python3 scope2026.py /private/path/C004-26003-01_Coates_Event_Labour_Scope.xlsx [previous_event_scope.json]

Writes event_scope_2026.json beside this script and prints the old-vs-new comparison. The workbook itself stays out of
the repository (rates); its SHA-256 is recorded in the JSON.
"""
import datetime as dt, hashlib, json, sys
from pathlib import Path
import openpyxl

HERE = Path(__file__).resolve().parent
src = Path(sys.argv[1]); prev = json.loads(Path(sys.argv[2]).read_text(encoding='utf-8')) if len(sys.argv) > 2 else None
wb = openpyxl.load_workbook(src, data_only=True); ws = wb['Sheet1']
sha = hashlib.sha256(src.read_bytes()).hexdigest()

def hhmm(v):
    if v is None: return None
    if isinstance(v, dt.time): return None if (v.hour == 0 and v.minute == 0) else f'{v.hour:02d}:{v.minute:02d}'
    raise ValueError(repr(v))

def hours_exact(v):
    if isinstance(v, dt.timedelta): return v.total_seconds() / 3600
    if isinstance(v, dt.time): return v.hour + v.minute / 60
    return float(v or 0)
def hours(v): return round(hours_exact(v), 4)

DAYS = [('Fri', '2026-10-23', 3), ('Sat', '2026-10-24', 4), ('Sun', '2026-10-25', 5)]
FIELDS = [('accredited_gates_open', 'Accredited gates open', 'B'), ('patron_gates_open', 'Patron gates open', 'C'), ('gate_closure', 'Gate closure commences', 'D'), ('track_opening', 'Track opening', 'E'), ('entertainment_concludes', 'Entertainment concludes', 'F')]
access_days = []
for day, date, row in DAYS:
    d = {'day': day, 'date': date, 'row': row}
    for key, label, col in FIELDS: d[key] = hhmm(ws[f'{col}{row}'].value)
    access_days.append(d)

# the lines: rate cell, the three day rows, the window as the workbook writes it, the invoicing codes, and the rules that
# give the written times from the access plan (the same rules as the 2025 transcription, except VMS: gate closure to
# track opening, as the 2026 times are written)
LINES = [
    dict(key='pm', section='day', role='Project Manager', card_role='Project Manager (per hr)', rate_cell='B8', rows=(10, 11, 12), rows_label='8–13',
         window='accredited gate opening to approx 30 min post track opening (Sunday approx 1.5 h)', invoicing='3.14 Portable Building Hire and 3.16 Toilet Hire, the hours split in half',
         days={'Fri': {'people': 1, 'start': {'at': 'accredited_gates_open'}, 'finish': {'at': 'track_opening', 'plus_min': 30, 'round': 'floor'}},
               'Sat': {'people': 1, 'start': {'at': 'accredited_gates_open'}, 'finish': {'at': 'track_opening', 'plus_min': 30, 'round': 'floor'}},
               'Sun': {'people': 1, 'start': {'at': 'accredited_gates_open'}, 'finish': {'at': 'track_opening', 'plus_min': 90, 'round': 'ceiling'}}}),
    dict(key='gentech', section='day', role='Gen Tech', card_role='Gen Tech (per hr)', rate_cell='B15', rows=(17, 18, 19), rows_label='15–20',
         window='30 min before patron gates to the conclusion of racing', invoicing='3.06 Electrical',
         days={d: {'people': 1, 'start': {'at': 'patron_gates_open', 'day': 'Fri', 'plus_min': -30, 'round': 'floor'}, 'finish': {'at': 'track_opening', 'plus_min': 15, 'round': 'floor'}} for d in ('Fri', 'Sat', 'Sun')}),
    dict(key='fence', section='day', role='Fence team', card_role='Labour - Fence Team (per hr)', rate_cell='B22', rows=(24, 25, 26), rows_label='22–27',
         window='15 min before gate closure to approx 15 min after track opening (Sunday approx 1.5 h)', invoicing='3.15 Security & Demarcation Fencing',
         days={'Fri': {'people': 6, 'start': {'at': 'gate_closure', 'plus_min': -15, 'round': 'ceiling'}, 'finish': {'at': 'track_opening', 'plus_min': 15, 'round': 'floor'}},
               'Sat': {'people': 6, 'start': {'at': 'gate_closure', 'plus_min': -15, 'round': 'floor'}, 'finish': {'at': 'track_opening', 'plus_min': 15, 'round': 'ceiling'}},
               'Sun': {'people': 6, 'start': {'at': 'gate_closure', 'plus_min': -15, 'round': 'floor'}, 'finish': {'at': 'track_opening', 'plus_min': 90, 'round': 'ceiling'}}}),
    dict(key='toilets', section='day', role='Toilets', card_role='Labour - Toilets (per hr)', rate_cell='B29', rows=(31, 32, 33), rows_label='29–34',
         window='30 min before patron gates to the conclusion of racing', invoicing='3.16 Toilet Hire',
         days={d: {'people': 1, 'start': {'at': 'patron_gates_open', 'day': 'Fri', 'plus_min': -30, 'round': 'floor'}, 'finish': {'at': 'track_opening', 'plus_min': 15, 'round': 'floor'}} for d in ('Fri', 'Sat', 'Sun')}),
    dict(key='vms', section='day', role='VMS board movements', card_role=None, rate_cell='B39', rows=(41, 42, 43), rows_label='39–44',
         window='gate closure to track opening, as the times are written (the note beside them still reads 30 min pre patron gates to the conclusion of racing)', invoicing='Traffic (no cost code on the line)',
         rate_from='the VMS section’s own rate cell (B39) — the card has no VMS role',
         days={d: {'people': 2, 'start': {'at': 'gate_closure'}, 'finish': {'at': 'track_opening'}} for d in ('Fri', 'Sat', 'Sun')}),
    dict(key='gentech_ent', section='entertainment', role='Gen Tech — entertainment', card_role='Gen Tech (per hr)', rate_cell='B49', rows=(51, 52, 53), rows_label='49–54',
         window="the day staff's shift end to the conclusion of the entertainment", invoicing='7.04 Concert',
         days={'Fri': {'people': 1, 'start': {'line': 'gentech'}, 'finish': {'at': 'entertainment_concludes'}, 'only_if': 'entertainment_concludes'},
               'Sat': {'people': 1, 'start': {'line': 'gentech'}, 'finish': {'at': 'entertainment_concludes'}, 'only_if': 'entertainment_concludes'},
               'Sun': {'people': 0, 'start': {'line': 'gentech'}, 'finish': {'at': 'entertainment_concludes'}, 'only_if': 'entertainment_concludes'}}),
    dict(key='toilets_ent', section='entertainment', role='Toilets — entertainment', card_role='Labour - Toilets (per hr)', rate_cell='B56', rows=(58, 59, 60), rows_label='56–61',
         window="the day staff's shift end to the conclusion of the entertainment", invoicing='7.04 Concert',
         days={'Fri': {'people': 1, 'start': {'line': 'toilets'}, 'finish': {'at': 'entertainment_concludes'}, 'only_if': 'entertainment_concludes'},
               'Sat': {'people': 1, 'start': {'line': 'toilets'}, 'finish': {'at': 'entertainment_concludes'}, 'only_if': 'entertainment_concludes'},
               'Sun': {'people': 0, 'start': {'line': 'toilets'}, 'finish': {'at': 'entertainment_concludes'}, 'only_if': 'entertainment_concludes'}}),
]

def mins(hm): return None if hm is None else int(hm[:2]) * 60 + int(hm[3:])
def r15(m, mode): return (m // 15) * 15 if mode == 'floor' else (-(-m // 15)) * 15 if mode == 'ceiling' else m
plan = {d['day']: d for d in access_days}

lines, finish_of = [], {}
for L in LINES:
    rate = float(ws[L['rate_cell']].value)
    aw = {'start': {}, 'finish': {}, 'hours': {}}
    total_hours = 0.0; amount = 0.0
    for (day, date, _), row in zip(DAYS, L['rows']):
        s, f, q = hhmm(ws[f'B{row}'].value), hhmm(ws[f'C{row}'].value), ws[f'D{row}'].value or 0
        hrs = hours(ws[f'F{row}'].value); hrs_exact = hours_exact(ws[f'F{row}'].value); cost = float(ws[f'G{row}'].value or 0)
        aw['start'][day] = s; aw['finish'][day] = f; aw['hours'][day] = hrs
        total_hours += hrs_exact; amount += cost
        # the rule reproduces the written times
        rule = L['days'][day]
        def at(spec):
            if 'line' in spec: return finish_of[spec['line'] + '|' + day]
            m = mins(plan[spec.get('day', day)][spec['at']])
            return None if m is None else r15(m + spec.get('plus_min', 0), spec.get('round'))
        on = ('only_if' not in rule) or plan[day][rule['only_if']] is not None
        rs, rf = (at(rule['start']), at(rule['finish'])) if on else (None, None)
        finish_of[L['key'] + '|' + day] = rf
        per = max(0, rf - rs) / 60 if rs is not None and rf is not None else 0
        got = round(per * rule['people'], 4)
        assert got == hrs, (L['key'], day, got, hrs, rs, rf)
        assert abs(rate * hrs_exact - cost) < 0.0005, (L['key'], day, rate * hrs_exact, cost)
        assert (rule['people'] if on else 0) == (q or 0) or (not on and not q), (L['key'], day, q)
        if rs is not None: assert f'{rs // 60:02d}:{rs % 60:02d}' == s and f'{rf // 60:02d}:{rf % 60:02d}' == f, (L['key'], day, s, f, rs, rf)
    aw['total_hours'] = round(total_hours, 4); aw['amount'] = round(amount, 4)
    line = {'key': L['key'], 'section': L['section'], 'role': L['role'], 'card_role': L['card_role'], 'rate': rate, 'rows': L['rows_label'], 'window': L['window'], 'invoicing': L['invoicing']}
    if L.get('rate_from'): line['rate_from'] = L['rate_from']
    line['days'] = L['days']; line['as_written'] = aw
    lines.append(line)

def cell(c): return round(float(ws[c].value), 4)
day_staffing = cell('G36'); vms_moves = cell('G44'); entertainment = cell('G63'); at_amount = cell('G68')
assert abs(sum(l['as_written']['amount'] for l in lines if l['section'] == 'day' and l['key'] != 'vms') - day_staffing) < 0.001
assert abs(next(l for l in lines if l['key'] == 'vms')['as_written']['amount'] - vms_moves) < 0.001
assert abs(sum(l['as_written']['amount'] for l in lines if l['section'] == 'entertainment') - entertainment) < 0.001
assert abs(cell('M68') + cell('M69') - at_amount) < 0.001 and abs(cell('M36') - day_staffing) < 0.001 and abs(cell('M63') - entertainment) < 0.001
people_total = round(sum(l['as_written']['amount'] for l in lines), 4)
grand = round(people_total + at_amount, 4)

scope = {
    'schema': 'gc500-event-labour-scope-1',
    'document': 'C004-26003-01 Coates Event Labour Scope',
    'file': 'C004-26003-01_Coates_Event_Labour_Scope.xlsx (held privately; not in the repository)',
    'sha256': sha,
    'supplied_by': 'Andrew Fisher', 'supplied_on': '2026-10-23',
    'sent_as': "please analyse these costs and update where we need to this is what we charge the v8's",
    'workbook': {'sheet': 'Sheet1', 'created': '2022-10-25 by Matt Koleczko', 'last_saved': '2026-10-23 12:30 by Fisher, Andrew (North)'},
    'is': "What Coates charges the V8s for people over the event: the staffing requirement on each race day and for the entertainment, worked out from the 2026 access plan's times, and accommodation and travel. Transcribed formula by formula from the workbook; every hour and dollar the page works out from it is checked against the workbook's own figures by the release's scope2026.py and tests/test_v1154.cjs.",
    'rounding': "a time rounded down (floor) or up (ceiling) is rounded to the quarter hour, as the workbook's FLOOR and CEILING do",
    'access_plan': {'fields': [{'key': k, 'label': lab, 'col': col} for k, lab, col in FIELDS], 'days': access_days,
                    'as_written': "the 2026 access plan as the workbook carries it on 23 Oct 2026: accredited and patron gates Fri 09:30, Sat and Sun 08:00; gate closure 07:00 each day; track opening 18:00, 18:00 and 17:40; entertainment concludes 20:30 Friday and Saturday, none Sunday",
                    'dates_from': 'the three race days on the record (Fri 23, Sat 24 and Sun 25 Oct 2026); the workbook names the days only'},
    'lines': lines,
    'accommodation_travel': {'amount': at_amount, 'row': 68, 'split': [{'cost_code': '3.14', 'name': 'Portable Building Hire', 'amount': cell('M68')}, {'cost_code': '3.16', 'name': 'Toilet Hire', 'amount': cell('M69')}],
                             'as_written': 'a figure typed in the workbook, charged half under Portable Building Hire and half under Toilet Hire'},
    'totals_as_written': {'day_staffing': day_staffing, 'vms_moves': vms_moves, 'entertainment': entertainment, 'accommodation_travel': at_amount, 'people': people_total, 'grand': grand},
    'read_notes': [
        "The rate cells are the workbook's own: $150.11 an hour for the Project Manager and Gen Tech, $91.14 for the fence team, toilets and VMS — 3% over the 2025 card roles the page's rate card lists ($145.74 and $88.49). The workbook's cells are what is charged.",
        "Gen Tech and Toilets start 30 minutes before Friday's patron-gate time on all three days — the workbook points Saturday's and Sunday's start at the Friday cell (C3): 09:00 every day, while Saturday's and Sunday's gates open at 08:00. Read as written.",
        "VMS board movements run gate closure (07:00) to track opening, as the times are written, at the VMS section's rate cell ($91.14); the note beside them still reads 30 min pre patron gates to the conclusion of racing.",
        "The VMS section's invoicing total (M46) adds up M20:M44, which takes in the day-staffing total (M36) and the fence and toilet lines a second time, and reads $58,453.63; the VMS lines themselves come to $5,954.48. The page adds up the lines, never that cell.",
        "The entertainment hours run from the day staff's finish (18:15) to the end of the entertainment (20:30) on Friday and Saturday. Sunday has no entertainment time and no one on the line, so none.",
    ],
}
out = HERE / 'event_scope_2026.json'
out.write_text(json.dumps(scope, indent=1, ensure_ascii=False) + '\n', encoding='utf-8')
print('written', out, '| workbook sha256', sha[:16], '| people', people_total, '| accommodation & travel', at_amount, '| grand', grand)
if prev:
    print('\nline            | 2025 doc hours / amount        | 2026 doc hours / amount')
    for new in lines:
        old = next((l for l in prev['lines'] if l['key'] == new['key']), None)
        o = old['as_written'] if old else {}
        print(f"{new['key']:15s} | {o.get('total_hours', 0):8.4f} h  {o.get('amount', 0):12.4f} @{old['rate'] if old else 0:7.2f} | {new['as_written']['total_hours']:8.4f} h  {new['as_written']['amount']:12.4f} @{new['rate']:7.2f}")
    ot = prev['totals_as_written']
    print(f"day staffing    | {ot['day_staffing']:12.4f} | {day_staffing:12.4f}\nvms             | {ot['vms_moves']:12.4f} | {vms_moves:12.4f}\nentertainment   | {ot['entertainment']:12.4f} | {entertainment:12.4f}\naccom & travel  | {ot['accommodation_travel']:12.4f} | {at_amount:12.4f}")
    print(f"people          | {sum(l['as_written']['amount'] for l in prev['lines']):12.4f} | {people_total:12.4f}\ngrand           | {sum(l['as_written']['amount'] for l in prev['lines']) + ot['accommodation_travel']:12.4f} | {grand:12.4f}")
