"""Author: Andrew Fisher. The toilets list for the branch (9968955): every reference's units on site (the servicing tracker,
Andrew's 13 areas) against Baseplan's hire lines, number by number, and the charge lines at the correct quantities.
    python3 -I build_branch_toilets.py <private dir>     (bp6.json in it; the page build beside)
"""
import json, re, sys, collections
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
HERE = Path(__file__).resolve().parent; PRIV = Path(sys.argv[1])
PAGE = HERE.parent / 'build' / 'GC500_v11.16' / 'GC500_Delivery_Control_hosted.html'
bp = json.load(open(PRIV / 'bp6.json'))
s = open(PAGE, encoding='utf-8').read(); i = s.find('var SVC1067_CARD='); j = s.find('\n', i); SVC = json.loads(s[i + len('var SVC1067_CARD='):j].rstrip(';'))['data']
KIND = {'fwf': 'FWF', 'access': 'Accessible', 'sewer': 'Block (sewer)', 'block': 'Block (tank-mounted)', 'tank': 'Waste tank', 'pan16': '16-pan block', 'pump_only': 'Pump-out only'}
trk = {}
for a in SVC['areas']:
    for r in a['refs']:
        t = trk.setdefault(r['ref'], {'area': f"{a['n']} {a['name']}", 'k': collections.Counter(), 'nos': collections.defaultdict(list)})
        for l in r['lines']:
            t['k'][l['kind']] += l['qty']; t['nos'][l['kind']] += [str(n) for n in (l.get('nos') or [])]
bpr = collections.defaultdict(lambda: {'k': collections.Counter(), 'nos': collections.defaultdict(list), 'lines': []})
for d in bp['9968955']:
    if str(d.get('Item Type')) != '1': continue
    desc = str(d.get('Description', '')); m = re.match(r'^(WC\d+\w*|PG\d+|Wc\d+)\s+(.*)', desc); head = (m.group(1) if m else '?').upper(); kind = (m.group(2) if m else desc)
    k = 'fwf' if ('Fresh Water' in kind or kind.startswith('FWF Toilet') or 'Trailer' in kind) else 'access' if ('Accessible' in kind or 'Disabled' in kind) else 'pan16' if '16 Pan' in kind else 'tank' if 'Holding Tank' in kind else 'blk' if 'Block' in kind else 'other'
    q = d.get('Quantity'); q = int(q) if isinstance(q, (int, float)) else 1
    bpr[head]['k'][k] += q; bpr[head]['lines'].append(str(d.get('Line')))
    if re.match(r'^\d{6,7}$', str(d.get('Item'))): bpr[head]['nos'][k].append(str(d.get('Item')))
wb = Workbook(); ws = wb.active; ws.title = '9968955 by reference'
ws.append(['Ref', 'Servicing area', 'Baseplan lines', 'FWF on site', 'FWF Baseplan', 'Accessible on site', 'Accessible Baseplan', 'Blocks on site (sewer + tank-mounted)', 'Blocks Baseplan', 'Tanks on site', 'Tanks Baseplan', '16-pan on site', '16-pan Baseplan', 'Coates numbers on site not on Baseplan', 'Baseplan numbers not on site', 'Fix needed on Baseplan'])
for c in ws[1]: c.font = Font(bold=True)
trk.pop('T0243', None)   # the WCTV toilet block 1311144 is contract 9973977, not 9968955
refs = sorted(set(trk) | set(bpr), key=lambda x: (x[:2], int(re.sub(r'\D', '', x) or 0), x))
T = collections.Counter(); fixes = 0
for ref in refs:
    t = trk.get(ref, {'area': '', 'k': collections.Counter(), 'nos': {}}); b = bpr.get(ref, {'k': collections.Counter(), 'nos': {}, 'lines': []})
    ofwf = t['k']['fwf']; bfwf = b['k']['fwf']; oacc = t['k']['access']; bacc = b['k']['access']; oblk = t['k']['sewer'] + t['k']['block']; bblk = b['k']['blk']
    onums = sorted(n for n in t['nos'].get('fwf', []) + t['nos'].get('access', []) + t['nos'].get('sewer', []) + t['nos'].get('block', []) + t['nos'].get('tank', []) + t['nos'].get('pan16', []) if re.match(r'^\d{6,7}$', n))
    bnums = sorted(n for ns in b['nos'].values() for n in ns)
    fix = []
    if ofwf != bfwf: fix.append(f"FWF {bfwf} -> {ofwf}")
    if oacc != bacc: fix.append(f"accessible {bacc} -> {oacc}")
    if oblk != bblk: fix.append(f"blocks {bblk} -> {oblk}")
    if t['k']['tank'] != b['k']['tank']: fix.append(f"tanks {b['k']['tank']} -> {t['k']['tank']}")
    if t['k']['pan16'] != b['k']['pan16']: fix.append(f"16-pan {b['k']['pan16']} -> {t['k']['pan16']}")
    miss_bp = sorted(set(onums) - set(bnums)); miss_site = sorted(set(bnums) - set(onums))
    if miss_site: fix.append('number(s) not on site: ' + ', '.join(miss_site))
    if miss_bp and bnums: fix.append('number(s) missing: ' + ', '.join(miss_bp))
    ws.append([ref, t['area'], ', '.join(b['lines']), ofwf, bfwf, oacc, bacc, oblk, bblk, t['k']['tank'], b['k']['tank'], t['k']['pan16'], b['k']['pan16'], ', '.join(miss_bp), ', '.join(miss_site), '; '.join(fix)])
    if fix: fixes += 1
    for k, v in (('ofwf', ofwf), ('bfwf', bfwf), ('oacc', oacc), ('bacc', bacc), ('oblk', oblk), ('bblk', bblk), ('otank', t['k']['tank']), ('btank', b['k']['tank']), ('opan', t['k']['pan16']), ('bpan', b['k']['pan16'])): T[k] += v
ws.append([]); ws.append(['TOTAL', '', '', T['ofwf'], T['bfwf'], T['oacc'], T['bacc'], T['oblk'], T['bblk'], T['otank'], T['btank'], T['opan'], T['bpan'], '', '', f"{fixes} references to fix"])
for c in ws[ws.max_row]: c.font = Font(bold=True)
for i, w in enumerate([7, 30, 16, 8, 8, 9, 9, 12, 9, 8, 8, 8, 8, 30, 24, 50]): ws.column_dimensions[chr(65 + i)].width = w
fill = PatternFill('solid', fgColor='F8CBAD')
for row in ws.iter_rows(min_row=2, max_row=ws.max_row - 2):
    if row[15].value:
        for c in row: c.fill = fill
ws.freeze_panes = 'A2'
# charge lines
svc = {str(d.get('Line')): d for d in bp['9968955'] if str(d.get('Item Type')) == '2'}
N = T['ofwf']; A = T['oacc']; BLK = T['oblk'] + T['opan']; SEW = T['oblk'] - 6; TK = T['otank'] + T['opan']; NIGHTS = 3
cs = wb.create_sheet('9968955 charge lines')
cs.append(['Line', 'Baseplan line', 'Baseplan qty', 'Rate', 'Baseplan $', 'Correct qty', 'Correct $', 'Billable now $', 'Later $', 'Basis'])
for c in cs[1]: c.font = Font(bold=True)
SUMS = collections.Counter()
def row(ln, qty, basis, now_frac):
    d = svc[ln]; amt = round(qty * d['Price'], 2); now = round(amt * now_frac, 2)
    cs.append([ln, d['Description'], d['Quantity'], d['Price'], d['Prebill Amount'], qty, amt, now, round(amt - now, 2), basis])
    SUMS['bp'] += d['Prebill Amount']; SUMS['ok'] += amt; SUMS['now'] += now; SUMS['later'] += amt - now
row('141', (N + A) * NIGHTS, f"{N} FWF + {A} accessible on site = {N + A} units x {NIGHTS} race nights (Fri, Sat, Sun). Nothing done yet; invoice from the tracker ticks after each night", 0)
row('142', TK * NIGHTS, f"{T['otank']} waste tanks + {T['opan']} sixteen-pan blocks = {TK} x {NIGHTS} nights; agrees with Baseplan", 0)
row('143', SEW * NIGHTS, f"{SEW} sewer-connected blocks x {NIGHTS} nights (the {T['oblk'] - SEW} tank-mounted blocks are cleaned under line 142); agrees with Baseplan", 0)
row('144', BLK, f"{T['oblk']} blocks + {T['opan']} sixteen-pan = {BLK}; $603.78 is install + steps + cleaning + demob. Install, steps and cleaning done ($458.04 each); demob ($145.74) at demob", 458.04 / 603.78)
row('145', A * 2, f"{A} accessible x 2 (install $88.49 done, demob $88.49 at demob)", 0.5)
row('146', N * 2, f"{N} FWF x 2 (install $36.44 done, demob $36.44 at demob)", 0.5)
row('147', BLK * 2, f"{BLK} blocks x 2 legs at $500; delivery done, return at demob", 0.5)
row('148', A * 2, f"{A} accessible x 2 legs at $500; delivery done, return at demob", 0.5)
row('149', N * 2, f"{N} FWF x 2 legs at $47.50; delivery done, return at demob. Baseplan already charges the Event Portables units this way", 0.5)
cs.append([]); cs.append(['TOTAL', '', '', '', round(SUMS['bp'], 2), '', round(SUMS['ok'], 2), round(SUMS['now'], 2), round(SUMS['later'], 2), ''])
for c in cs[cs.max_row]: c.font = Font(bold=True)
for i, w in enumerate([6, 40, 10, 8, 12, 10, 12, 13, 11, 90]): cs.column_dimensions[chr(65 + i)].width = w
cs.freeze_panes = 'A2'
n = wb.create_sheet('Notes')
for line in ['Author: Andrew Fisher. 22 Oct 2026. Units on site are the servicing tracker (Andrew\'s 13 areas, read 19-20 Oct from the site), one row a reference.',
             'Rates: 2026 card, ex GST. Pump-out FWF/accessible $72.87 a service; tank pump-out $624.60; sewer-connect clean & restock $260.25; block install & demob $603.78; accessible $88.49 a leg; FWF $36.44 a leg; transport $500 (blocks, accessible) and $47.50 (FWF) each way.',
             'Services are the three race nights (Fri 23, Sat 24, Sun 25 Oct). None has happened yet: invoice them from the tracker ticks after each night, never in advance.',
             'Rows shaded red need a Baseplan fix (named in the last column).']: n.append([line])
n.column_dimensions['A'].width = 160
wb.save(HERE / 'GC500_9968955_Toilets_for_Branch_22Oct2026.xlsx')
print({k: v for k, v in T.items()}, 'fix refs', fixes, 'N', N, 'A', A, 'BLK', BLK, 'SEW', SEW, 'TK', TK)
for r in cs.iter_rows(min_row=2, values_only=True): print(r[:9])
