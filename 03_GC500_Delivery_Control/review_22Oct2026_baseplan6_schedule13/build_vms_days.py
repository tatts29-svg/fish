"""Author: Andrew Fisher. 9961265 VMS boards, day by day: a column a day from the first board landing (7 Sep) to the last race day
(25 Oct), the card's daily rate in each cell from the day the board went in (VMS is charged from delivery), with install,
demob, the delivery and pickup legs, and a map link to each board's position as the page holds it (the pin taken on site
where there is one, else the plan).
    python3 -I build_vms_days.py <private dir>      (vms_model.json - the page's own VMS model, read read-only - and bp6.json in it)
"""
import json, sys, datetime as dt, collections
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as col
HERE = Path(__file__).resolve().parent; PRIV = Path(sys.argv[1])
M = json.load(open(PRIV / 'vms_model.json')); bp = json.load(open(PRIV / 'bp6.json'))
RATE = 53.6115; LAB = 88.485; LEG = 150.0; MIN_DAYS = 7
D0, D1 = dt.date(2026, 9, 7), dt.date(2026, 10, 25); TODAY = dt.date(2026, 10, 22); RACE = {dt.date(2026, 10, 23), dt.date(2026, 10, 24), dt.date(2026, 10, 25)}
days = [D0 + dt.timedelta(n) for n in range((D1 - D0).days + 1)]
LOADS = {'T0001': '2026-09-07', 'T0103': '2026-10-08', 'T0128': '2026-10-16', 'T0158': '2026-10-19'}
WORD = {'VMS22': '2026-10-22', 'VMS24': '2026-10-22'}   # Andrew, 22 Oct 2026: today's delivery, docket 26143911 (the record said 21 Oct)
def aest(iso): return (dt.datetime.fromisoformat(iso.replace('Z', '+00:00')) + dt.timedelta(hours=10)).date().isoformat()
pos = {p['pos']: p for p in M['positions']}
b913 = {b['key']: b for b in M['boards913']}
# the board each position holds now; a board that moved to its event position is shown there
where = {}
for p in M['positions']:
    key = p['board']
    if p['pos'].endswith('a'):
        r = p.get('position') or {}
        if 'pinned' in str(r.get('how', '')) or 'confirmed' in str(r.get('how', '')) or 'GPS' in str(r.get('how', '')): where[key] = p   # moved: the pin on site proves it
        continue
    where.setdefault(key, p)
for p in M['positions']:
    if p['pos'].endswith('a') and where.get(p['board'], {}).get('pos') == p['pos']: continue
    if p['pos'].endswith('a'): continue
    where.setdefault(p['board'], p)
# a board named for its position sits there (VMS09 is at 09; the plan also lists 1211404 against 05, where it stood 7-22 Sep)
for key in list(where):
    if key.startswith('VMS') and key[3:] in pos and pos[key[3:]]['board'] == key and not where[key]['pos'].endswith('a'): where[key] = pos[key[3:]]
def on_site(key, p):
    """the day the board went in, and the basis: the record's arrival event; Andrew's word of 22 Oct for VMS22 and 24; the
    page's own 'on' reading (a date he said, a load he recorded it on, or the contract line's 7 Sep)"""
    if key in WORD: return WORD[key], "Andrew, 22 Oct: today's delivery (docket 26143911); the record's 21 Oct is to be corrected"
    a = p.get('arrival')
    if a and a.get('effective_date'): return a['effective_date'], 'arrival recorded on the page by ' + str(a.get('by', '')).split(' via ')[0]
    if key == 'VMS05': return '2026-10-20', "pinned at 05a by Andrew Fisher 20 Oct 10:08 AEST; docket 26141143 the same morning"
    o = (b913.get(key) or {}).get('on') or {}
    if o.get('src') == 'word' and o.get('said'):
        d = dt.datetime.strptime(o['said'], '%d %b %Y').date().isoformat(); return d, f"Andrew's word on the page: {o['said']}"
    if o.get('src') == 'record' and o.get('at'):
        d = aest(o['at']); return d, f"recorded on site {d} by {str(o.get('by', '')).split(' via ')[0]} (Timeline load {o.get('on')}, {LOADS.get(o.get('on'), '?')})"
    if o.get('src') == 'contract' and o.get('on') in LOADS: return LOADS[o['on']], 'contract line and docket, 7 Sep (Timeline load T0001)'
    return None, 'no on-site day on the record'
bpl = {}
for d in bp['9961265']:
    if str(d.get('Item Type')) != '1' or d.get('Status') == 'Returned': continue
    desc = str(d.get('Description', '')); key = desc.split()[0] if desc.startswith('VMS') else str(d.get('Item'))
    if key == 'VMS05a': key = 'VMS05'
    bpl[key] = d
order = sorted(where, key=lambda k: (int(''.join(c for c in (b913.get(k, {}).get('name') or k) if c.isdigit()) or 0)))
wb = Workbook(); ws = wb.active; ws.title = 'Day by day'
H = Font(bold=True); ON = PatternFill('solid', fgColor='DDEBF7'); FULL = PatternFill('solid', fgColor='FCE4D6'); TOT = PatternFill('solid', fgColor='E2EFDA'); LINK = Font(color='0563C1', underline='single')
fixed = ['Board', 'Asset / fleet', 'Rego', 'Owner', 'Position now', 'On site', 'Map']
ws.append(fixed + [('RACE' if d in RACE else '') for d in days] + ['Days to 22 Oct', 'Hire to 22 Oct', 'Race days', 'Hire 23-25 Oct', 'Hire to 25 Oct (min 7)', 'Install (done)', 'Demob (later)', 'Delivery leg (done)', 'Pickup leg (later)', 'Billable now', 'Later', 'Board total to 25 Oct', 'Baseplan line', 'Baseplan start', 'Baseplan prebill', 'Basis / note'])
ws.append([''] * len(fixed) + [d.strftime('%a %d %b') for d in days] + [''] * 16)
for r in (1, 2):
    for c in ws[r]: c.font = H; c.alignment = Alignment(horizontal='center', wrap_text=True)
daytot = [0.0] * len(days); G = collections.Counter(); rows_out = []
for key in order:
    p = where[key]; b = b913.get(key, {}); idn = p.get('identity') or {}; P = p.get('position') or {}
    name = b.get('name') or key; on, basis = on_site(key, p)
    lat, lon = P.get('lat'), P.get('lon'); link = f'https://www.google.com/maps?q={lat:.6f},{lon:.6f}' if lat is not None else ''
    cells = []; d22 = 0; h22 = 0.0; rd = 0; hr = 0.0
    start = dt.date.fromisoformat(on) if on else None
    for n, d in enumerate(days):
        if not start or d < start: cells.append(None); continue
        cells.append(round(RATE, 2)); daytot[n] += RATE
        if d in RACE: rd += 1; hr += RATE
        else: d22 += 1; h22 += RATE
    h25 = max(d22 + rd, MIN_DAYS) * RATE if start else 0
    now = h22 + LAB + LEG; later = LAB + LEG; total = h25 + 2 * LAB + 2 * LEG
    L = bpl.get(key) or bpl.get(str(idn.get('asset'))) or {}
    note = basis
    if L and on and str(L.get('Start Date'))[:10] != on: note += f"; Baseplan starts it {str(L.get('Start Date'))[:10]}"
    ws.append([name, idn.get('asset') or '', idn.get('rego') or '', idn.get('co') or '', f"VMS {p['pos'].upper().replace('A', 'a')} · {(p.get('title') or '')[:60]}", start.strftime('%d %b') if start else '', 'Map'] + cells +
              [d22, round(h22, 2), rd, round(hr, 2), round(h25, 2), LAB, LAB, LEG, LEG, round(now, 2), round(later, 2), round(total, 2), L.get('Line', ''), str(L.get('Start Date', ''))[:10], L.get('Prebill Amount', ''), note])
    c = ws.cell(row=ws.max_row, column=len(fixed));
    if link: c.hyperlink = link; c.font = LINK
    rows_out.append((name, idn.get('asset'), idn.get('rego'), p['pos'], start, d22, round(h22, 2), round(h25, 2), round(now, 2), round(later, 2), round(total, 2), link))
    G['d22'] += d22; G['h22'] += h22; G['rd'] += rd; G['hr'] += hr; G['h25'] += h25; G['inst'] += LAB; G['dem'] += LAB; G['del'] += LEG; G['pick'] += LEG; G['now'] += now; G['later'] += later; G['tot'] += total; G['bp'] += L.get('Prebill Amount', 0) or 0
ws.append(['TOTAL', f'{len(order)} boards'] + [''] * (len(fixed) - 2) + [round(v, 2) for v in daytot] + [G['d22'], round(G['h22'], 2), G['rd'], round(G['hr'], 2), round(G['h25'], 2), round(G['inst'], 2), round(G['dem'], 2), G['del'], G['pick'], round(G['now'], 2), round(G['later'], 2), round(G['tot'], 2), '', '', round(G['bp'], 2), 'Baseplan prebills each line to its expected term (13 Nov for the 7 Sep boards)'])
for c in ws[ws.max_row]: c.font = H; c.fill = TOT
f0 = len(fixed) + 1
for row in ws.iter_rows(min_row=3, max_row=ws.max_row - 1):
    for n, d in enumerate(days):
        c = row[f0 - 1 + n]
        if c.value is not None: c.fill = FULL if d in RACE else ON; c.number_format = '#,##0.00'
    for k in range(f0 + len(days), f0 + len(days) + 15):
        if isinstance(row[k - 1].value, float): row[k - 1].number_format = '#,##0.00'
ws.freeze_panes = ws.cell(row=3, column=f0)
for i, w in enumerate([7, 10, 8, 12, 44, 7, 6]): ws.column_dimensions[col(i + 1)].width = w
for n in range(len(days)): ws.column_dimensions[col(f0 + n)].width = 7.5
for k, w in enumerate([7, 10, 6, 9, 11, 8, 8, 9, 9, 10, 8, 11, 7, 10, 10, 70]): ws.column_dimensions[col(f0 + len(days) + k)].width = w
ws.row_dimensions[2].height = 30
# the same boards, charges and map on one plain sheet
cs = wb.create_sheet('Charges & map')
cs.append(['Board', 'Asset / fleet', 'Rego', 'Owner', 'Position now', 'Where', 'Lat', 'Lon', 'Map link', 'On site', 'Days to 22 Oct', 'Hire to 22 Oct', 'Hire 23-25 Oct', 'Hire to 25 Oct (min 7)', 'Install (done)', 'Demob (later)', 'Delivery leg (done)', 'Pickup leg (later)', 'Billable now', 'Later', 'Board total to 25 Oct', 'Basis / note'])
for c in cs[1]: c.font = H; c.alignment = Alignment(wrap_text=True, horizontal='center')
for key in order:
    p = where[key]; b = b913.get(key, {}); idn = p.get('identity') or {}; P = p.get('position') or {}; on, basis = on_site(key, p)
    lat, lon = P.get('lat'), P.get('lon'); link = f'https://www.google.com/maps?q={lat:.6f},{lon:.6f}' if lat is not None else ''
    start = dt.date.fromisoformat(on) if on else None
    d22 = (TODAY - start).days + 1 if start else 0; rd = 3 if start else 0; h22 = d22 * RATE; hr = rd * RATE; h25 = max(d22 + rd, MIN_DAYS) * RATE if start else 0
    cs.append([b.get('name') or key, idn.get('asset') or '', idn.get('rego') or '', idn.get('co') or '', 'VMS ' + p['pos'].upper().replace('A', 'a'), (p.get('title') or ''), lat, lon, link, start.strftime('%d %b') if start else '', d22, round(h22, 2), round(hr, 2), round(h25, 2), LAB, LAB, LEG, LEG, round(h22 + LAB + LEG, 2), round(LAB + LEG, 2), round(h25 + 2 * LAB + 2 * LEG, 2), basis + ('' if 'pinned' in str(P.get('how', '')) or 'confirmed' in str(P.get('how', '')) or 'GPS' in str(P.get('how', '')) else ' · position is the plan, not a pin taken on site')])
    c = cs.cell(row=cs.max_row, column=9)
    if link: c.hyperlink = link; c.font = LINK
cs.append([]); cs.append(['TOTAL', f'{len(order)} boards', '', '', '', '', '', '', '', '', G['d22'], round(G['h22'], 2), round(G['hr'], 2), round(G['h25'], 2), round(G['inst'], 2), round(G['dem'], 2), G['del'], G['pick'], round(G['now'], 2), round(G['later'], 2), round(G['tot'], 2), ''])
for c in cs[cs.max_row]: c.font = H; c.fill = TOT
for i, w in enumerate([7, 10, 8, 12, 9, 50, 11, 11, 44, 7, 7, 10, 10, 11, 8, 8, 9, 9, 10, 8, 11, 70]): cs.column_dimensions[col(i + 1)].width = w
for rw in cs.iter_rows(min_row=2, max_row=cs.max_row):
    for c in rw:
        if isinstance(c.value, float) and c.column not in (7, 8): c.number_format = '#,##0.00'
cs.freeze_panes = 'B2'
n = wb.create_sheet('Rates & basis')
for line in ['Author: Andrew Fisher. 22 Oct 2026. Positions, identities and arrivals are the page\'s own VMS model, read read-only: the pin taken on site where Andrew took one (20-22 Oct), else the plan position; the map link opens that spot.',
             'Rate: 2026 street card, VMS Board, Daily Rate 2026 $53.6115, ex GST, on the traffic-management block that carries a 7-day minimum. VMS is charged from the day the board goes in (the project manager\'s rule of 17 Sep 2026) until it comes off at demob.',
             'Labour: the card\'s one VMS figure $176.97 (headed 2024), read as $88.49 install + $88.49 demob; install is done on every board on site, demob is at demob. Transport: the card says POA; $150 a leg is Baseplan\'s figure on this contract, pending Andrew\'s word.',
             'Hire to 22 Oct is what has run. 23-25 Oct are the race days (shaded), shown but not yet run. Boards stay on hire until demob, so "to 25 Oct" is not the final hire.',
             'On-site day: the record\'s arrival event (VMS21); Andrew\'s word of 22 Oct for VMS22 and VMS24 (today\'s delivery, docket 26143911; the record says 21 Oct and is to be corrected); the date Andrew said (VMS09, VMS10: 8 Oct); the day he recorded the board on site (VMS11-20, 23); the contract and docket for the seven 7 Sep boards; VMS05 the pin taken at 05a on 20 Oct.']: n.append([line])
n.column_dimensions['A'].width = 180
out = HERE / 'GC500_9961265_VMS_DayByDay_22Oct2026.xlsx'; wb.save(out)
print('written', out.name, {k: round(v, 2) for k, v in G.items()})
for r in rows_out: print(r[:11])
