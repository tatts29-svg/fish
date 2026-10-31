"""Author: Andrew Fisher. Published-source and exact-change boundary tests."""
import importlib.util
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('patch_v1250', ROOT / 'patch_v1250.py')
mod = importlib.util.module_from_spec(spec)
spec.loader.exec_module(mod)


class QuantityTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.base = Path(os.environ['BASE']).read_text(encoding='utf-8')
        assert mod.digest(cls.base.encode()) == mod.REVIEWED_BASE_SHA256
        cls.source = mod.source()

    def call(self, text=None, **kw):
        return mod.patch(self.base if text is None else text,
                         expected_base_sha=kw.get('sha', mod.REVIEWED_BASE_SHA256),
                         expected_base_version=kw.get('version', mod.REVIEWED_BASE_VERSION))

    def test_public_workbook_cells_and_identity(self):
        import openpyxl
        file = Path(os.environ['REVIEW_WORKBOOK'])
        self.assertEqual(mod.digest(file.read_bytes()), self.source['source']['sha256'])
        book = openpyxl.load_workbook(file, data_only=True)
        sheet = book[self.source['source']['sheet']]
        for rule in self.source['rows']:
            self.assertEqual(str(sheet[rule['source_contract_cell']].value), rule['contract'])
            self.assertEqual(sheet[rule['source_line_cell']].value, rule['line'])
            self.assertEqual(sheet[rule['source_quantity_cell']].value, rule['quantity'])
            row = sheet[rule['source_quantity_cell']].row
            self.assertEqual(sheet.cell(row, 5).value, rule['description'])
            self.assertEqual(sheet.cell(row, 4).value, rule['item'])
            self.assertEqual(sheet.cell(row, 9).value, rule['sales_code'])

    def test_only_six_quantities_and_footer_change(self):
        candidate = self.call()
        before, after = mod.source_rows(self.base), mod.source_rows(candidate)
        rules = {(x['contract'], x['line']): x for x in self.source['rows']}
        changed = []
        self.assertEqual(len(before), len(after))
        for old, new in zip(before, after):
            key = (str(old['rental_contract']), old['line'])
            if key in rules:
                self.assertEqual(new, dict(old, quantity=rules[key]['quantity']))
                changed.append(key)
            else:
                self.assertEqual(old, new)
        self.assertEqual(set(changed), set(rules))
        self.assertEqual(mod.reverse(candidate), self.base)
        self.assertEqual(candidate.count(mod.NEW_FOOTER), 1)
        self.assertNotIn(mod.OLD_FOOTER, candidate)

    def test_missing_explicit_inputs_rejected(self):
        with patch.dict(os.environ, {}, clear=True):
            with self.assertRaises(ValueError):
                mod.patch(self.base)

    def test_changed_predecessor_or_wrong_hash_rejected(self):
        for text, sha in [(self.base + '\n', mod.REVIEWED_BASE_SHA256), (self.base, '0' * 64)]:
            with self.assertRaises(ValueError):
                self.call(text, sha=sha)

    def test_wrong_version_and_reapplication_rejected(self):
        with self.assertRaises(ValueError):
            self.call(version='12.48')
        with self.assertRaises(ValueError):
            self.call(self.call())

    def test_manifest_change_rejected(self):
        with tempfile.TemporaryDirectory() as name:
            folder = Path(name)
            (folder / 'contract_quantity1250.json').write_text(json.dumps(self.source) + ' ')
            with patch.object(mod, 'ROOT', folder):
                with self.assertRaises(ValueError):
                    self.call()

    def test_identity_quantity_and_asset_guards(self):
        original = mod.source_rows(self.base)
        rule = self.source['rows'][0]
        row = next(r for r in original if str(r['rental_contract']) == rule['contract'] and r['line'] == rule['line'])
        encode = lambda r: json.dumps(r, ensure_ascii=False, separators=(',', ':'))
        for change in [{'quantity': 99}, {'description': 'Different source'}, {'branch_code': 'MEAD'},
                       {'charge_line': True}, {'start_date': '2026-10-25'},
                       {'match': dict(row['match'], key='P01')}]:
            modified = self.base.replace(encode(row), encode(dict(row, **change)), 1)
            with self.assertRaises(ValueError):
                mod.change_rows(modified)

    def test_duplicate_missing_and_changed_anchors_rejected(self):
        for modified in [self.base + '\nconst DATA = {}',
                         self.base.replace(mod.DATA_ANCHOR, 'const OTHER = ', 1)]:
            with self.assertRaises(ValueError):
                mod.source_rows(modified)
        for modified in [self.base.replace(mod.OLD_FOOTER, "+' · wrong';", 1),
                         self.base.replace('const previousContractCharge1249=', 'const changed=', 1)]:
            sha = mod.digest(modified.encode())
            with patch.object(mod, 'REVIEWED_BASE_SHA256', sha):
                with self.assertRaises(ValueError):
                    self.call(modified, sha=sha)

    def test_reverse_rejects_unrelated_change(self):
        with self.assertRaises(ValueError):
            mod.reverse(self.call() + '\n')

    def test_failed_cli_keeps_page_bytes(self):
        with tempfile.TemporaryDirectory() as name:
            file = Path(name) / 'page.html'
            file.write_text(self.base)
            env = dict(os.environ, GC500_V1250_BASE_SHA256='0' * 64,
                       GC500_V1250_BASE_VERSION=mod.REVIEWED_BASE_VERSION)
            result = subprocess.run([sys.executable, str(ROOT / 'patch_v1250.py'), str(file)],
                                    env=env, capture_output=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(file.read_text(), self.base)


if __name__ == '__main__':
    unittest.main()
