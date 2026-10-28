"""Synthetic source-boundary checks. Author: Andrew Fisher."""
import copy
import hashlib
import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

MODULE = Path(__file__).resolve().parents[1] / "source_adapter.py"
spec = importlib.util.spec_from_file_location("source_adapter1254", MODULE)
adapter = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = adapter
spec.loader.exec_module(adapter)


class ScopedSourceTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / "synthetic.xlsx"
        self.book = openpyxl.Workbook()
        self.book.remove(self.book.active)
        for contract, branch in adapter.CONTRACTS.items():
            sheet = self.book.create_sheet(contract + "-" + branch)
            sheet.append(adapter.HEADERS)
            raw = {"Line": 1, "Status": "Delivered", "Item Type": 3, "Item": "LABOUR", "Description": "Synthetic installation",
                   "Quantity": 1.25, "Price": 0, "Rate 1": None, "Sales Analysis Code": branch + "-LAB"}
            sheet.append([raw.get(k) for k in adapter.HEADERS])
        self.book.create_sheet("9999999-KINP").append(["deliberately unruled"])

    def tearDown(self):
        self.book.close()
        self.tmp.cleanup()

    def read(self):
        self.book.save(self.path)
        digest = hashlib.sha256(self.path.read_bytes()).hexdigest()
        return adapter.read_scoped_export(self.path, digest, dict(adapter.CONTRACTS))

    def data(self, source):
        rows = []
        for source_row in source.rows:
            row = copy.deepcopy(source_row)
            row.update(price=99, quantity=9, match={"state": "same asset number", "to": "asset", "key": "SYNTH", "via": "record"},
                       family="installation", kind="installation", charge_line=True, business_note="preserve")
            rows.append(row)
        rows.insert(1, {"rental_contract": "9999999", "branch_code": "KINP", "line": 7, "price": 88, "match": {"key": "UNCHANGED"}})
        return {"rental_on_hire": {"rows": rows, "assignments": {"SYNTH": {"lines": ["synthetic"]}}, "contracts": [{"keep": True}], "summary": {"keep": True}},
                "assets": [{"asset_numbers": ["1234567"]}], "ops": {"history": ["keep"]}, "plant_lines": {"lines": [{"rental": ["keep"]}]}}

    def test_zero_blank_decimal_and_item_type_survive(self):
        source = self.read()
        for row in source.rows:
            self.assertEqual(row["quantity"], 1.25)
            self.assertEqual(row["price"], 0)
            self.assertIsNone(row["rate_1"])
            self.assertEqual(row["item_type_code"], 3)
            self.assertIs(type(row["item_type_code"]), int)

    def test_fingerprint_requires_full_exact_sha(self):
        self.read()
        for sha in ["a" * 64, "12345678", None]:
            with self.subTest(sha=sha), self.assertRaises(adapter.SourceError):
                adapter.read_scoped_export(self.path, sha, dict(adapter.CONTRACTS))

    def test_explicit_scope_can_refresh_subset_only(self):
        self.book.save(self.path)
        digest = hashlib.sha256(self.path.read_bytes()).hexdigest()
        source = adapter.read_scoped_export(self.path, digest, {"9973977": "KINP"})
        self.assertEqual(len(source.rows), 1)
        self.assertEqual(source.approved_contracts, (("9973977", "KINP"),))
        for scope in [{}, {"9999999": "KINP"}, {"9973977": "MEAD"}]:
            with self.subTest(scope=scope), self.assertRaises(adapter.SourceError):
                adapter.read_scoped_export(self.path, digest, scope)

    def test_known_numeric_first_header_is_normalised(self):
        self.book["9968862-KINP"].cell(1, 1, 9968862)
        self.assertEqual(len(self.read().rows), 5)

    def test_missing_sheet_rejected(self):
        self.book.remove(self.book["9987005-MEAD"])
        with self.assertRaises(adapter.SourceError): self.read()

    def test_schema_rejected_even_when_price_column_moves(self):
        sheet = self.book["9987005-MEAD"]
        sheet.cell(1, 8, "Rate 1")
        with self.assertRaises(adapter.SourceError): self.read()

    def test_duplicate_identity_rejected(self):
        sheet = self.book["9987005-MEAD"]
        sheet.append([c.value for c in sheet[2]])
        with self.assertRaises(adapter.SourceError): self.read()

    def test_item_type_requires_numeric_integer(self):
        sheet = self.book["9987005-MEAD"]
        for value in [None, True, "2", 1.5, -1]:
            sheet.cell(2, 3).value = value
            with self.subTest(value=value), self.assertRaises(adapter.SourceError): self.read()

    def test_branch_rejected(self):
        self.book["9987005-MEAD"].title = "9987005-KINP"
        with self.assertRaises(adapter.SourceError): self.read()

    def test_ambiguous_quantities_and_formulas_rejected(self):
        sheet = self.book["9987005-MEAD"]
        for value in [None, True, "1.25", -1, "=1+1"]:
            sheet.cell(2, 7).value = value
            with self.subTest(value=value), self.assertRaises(adapter.SourceError): self.read()

    def test_scoped_merge_preserves_unruled_and_operational_data(self):
        source = self.read()
        data = self.data(source)
        before = copy.deepcopy(data)
        result, manifest = adapter.refresh_scoped_rows(data, source)
        self.assertEqual(data, before)
        self.assertEqual(result["rental_on_hire"]["rows"][1], before["rental_on_hire"]["rows"][1])
        for key in ("assignments", "contracts", "summary"):
            self.assertEqual(result["rental_on_hire"][key], before["rental_on_hire"][key])
        for key in ("assets", "ops", "plant_lines"):
            self.assertEqual(result[key], before[key])
        for row, old in zip(result["rental_on_hire"]["rows"], before["rental_on_hire"]["rows"]):
            self.assertEqual(row["match"], old["match"])
            if row["rental_contract"] in adapter.CONTRACTS:
                self.assertEqual(row["price"], 0)
                self.assertEqual(row["quantity"], 1.25)
                self.assertEqual(row["business_note"], "preserve")
        self.assertEqual(len(manifest["changed"]), 5)

    def test_removal_rejected(self):
        source = self.read()
        data = self.data(source)
        row = copy.deepcopy(data["rental_on_hire"]["rows"][0]);row["line"] = 2
        data["rental_on_hire"]["rows"].append(row)
        with self.assertRaises(adapter.SourceError): adapter.refresh_scoped_rows(data, source)

    def test_bundle_tampering_rejected(self):
        source = self.read();source.rows[0]["price"] = 999
        with self.assertRaises(adapter.SourceError): adapter.refresh_scoped_rows(self.data(source), source)

    def test_existing_rate_override_rejected(self):
        source = self.read();data = self.data(source)
        data["rental_on_hire"]["rows"][0]["rate_held1029"] = True
        with self.assertRaises(adapter.SourceError): adapter.refresh_scoped_rows(data, source)

    def test_new_rows_require_unmatched_explicit_template(self):
        source = self.read();data = self.data(source)
        removed = data["rental_on_hire"]["rows"].pop(0);key = (removed["rental_contract"], removed["line"])
        with self.assertRaises(adapter.SourceError): adapter.refresh_scoped_rows(data, source)
        with self.assertRaises(adapter.SourceError): adapter.refresh_scoped_rows(data, source, {key: removed})
        removed["match"] = {"state": "unmatched", "to": None, "key": None, "task_id": None, "via": None}
        result, manifest = adapter.refresh_scoped_rows(data, source, {key: removed})
        self.assertEqual(result["rental_on_hire"]["rows"][-1]["match"], removed["match"])
        self.assertEqual(sum(x["added"] for x in manifest["changed"]), 1)


if __name__ == "__main__":
    unittest.main()
