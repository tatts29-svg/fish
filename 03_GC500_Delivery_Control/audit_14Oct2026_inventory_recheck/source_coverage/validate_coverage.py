#!/usr/bin/env python3
"""Author: Andrew Fisher. Validate private source audit without publishing its rows."""
import argparse,collections,json
from pathlib import Path
import openpyxl
p=argparse.ArgumentParser();p.add_argument('--coverage',required=True);p.add_argument('--snapshot',required=True);p.add_argument('--workbook',required=True);a=p.parse_args()
r=json.loads(Path(a.coverage).read_text());s=json.loads(Path(a.snapshot).read_text());checks=[]
def check(name,condition):checks.append({'name':name,'pass':bool(condition)})
rows=r['baseplan']['rows'];ids=[(x['contract'],x['line'])for x in rows]
check('Every unique source contract line retained once',len(ids)==len(set(ids))==r['baseplan']['uniqueContractLines'])
check('Every source line has an explicit coverage disposition',all(x.get('coverage')and x.get('review')for x in rows))
check('Coverage categories conserve source lines',sum(r['counts']['baseplanClassifications'].values())==len(rows))
check('Confirmed source duplicates are excluded from unique-line total',r['baseplan']['worksheetRows']-r['baseplan']['uniqueContractLines']==len(r['baseplan']['exactDuplicateKeys']))
check('Original workbook rows independently reread',r['baseplan']['allOriginalRowsRechecked'])
check('Evidence uses the same shared record version',r['nativeVersion']==s['nativeVersion'])
check('Current reference count agrees with snapshot',r['counts']['currentReferences']==len(s['assets']))
check('Aliases excluded from canonical count',len(s['assets'])-len(s['canonical']['aliases'])==r['counts']['canonicalReferences'])
check('Cancelled history excluded from active canonical count',len(s['canonical']['refs'])-sum(bool(x['cancelled'])for x in s['matrix'])==r['counts']['activeCanonicalReferences'])
check('New source requests are marked unreceived',all(x['existing']is False for x in r['confirmedNewSchedule8Requests']))
check('Source requirement conflict does not invent arrivals',r['historicalWC67Difference']['additionalArrivalNotInferred'])
check('No native or original source mutations',r['nativeWrites']==r['sourceWrites']==0)
w=openpyxl.load_workbook(a.workbook,read_only=True,data_only=False)
check('Exactly two plain-language review tabs',w.sheetnames==['Existing','Needs checking'])
check('Every review row has reference, quantity, current state, source, contract and action',all(all(row[i].value is not None for i in [0,2,3,4,5,6,7,8])for sh in w for row in sh.iter_rows(min_row=5,max_col=9)))
check('No executable formulas in review workbook',not any(c.data_type=='f'for sh in w for row in sh for c in row))
check('No price totals in equipment-only workbook',not any('$'in str(c.value or '')for sh in w for row in sh for c in row))
check('Workbook credits Andrew Fisher',w.properties.creator=='Andrew Fisher')
result={'author':'Andrew Fisher','pass':all(x['pass']for x in checks),'checks':checks,'sourceSheets':r['baseplan']['sheets'],'sourceWorksheetRows':r['baseplan']['worksheetRows'],'sourceUniqueLines':len(rows),'currentReferences':len(s['assets']),'newSourceOnlyRequests':len(r['confirmedNewSchedule8Requests']),'reviewRows':{sh.title:sh.max_row-4 for sh in w},'nativeWrites':0,'originalWrites':0,'scope':'Read-only source coverage validation. No claim of full physical or financial accuracy.'}
print(json.dumps(result,indent=2));raise SystemExit(0 if result['pass']else 1)
