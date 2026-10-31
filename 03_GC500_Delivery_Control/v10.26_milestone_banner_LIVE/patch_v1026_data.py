"""Author: Andrew Fisher. v10.26 data: the Baseplan SuperCars export of 15 Oct 2026 (13 contract tabs, 326 lines)
refreshes the page's contract lines with v8.71's rules through the v9.49 reader (baseplan1026.py): rates, statuses,
dates, item numbers and docket numbers as exported; lines the export no longer has go to the source history; lines it
adds come in with their family rule; a line whose item number the export filled in re-joins its reference by that
number. Three joins are set by hand from the branch's delivery dockets (the Delivery Dockets sheet of the same export).
Schedule 9 (iEDM, 15 Oct 2026) then moves the dated events and quantities the page holds from Schedule 7, with the
site record kept authoritative: a reference already on site keeps its first day. Two review blocks record it all.

    BASEPLAN1026_XLSX=/private/Baseplan.xlsx SCHEDULE9_XLSX=/private/Schedule_9.xlsx RECORD_JSON=/private/state.json \
      python3 patch_v1026_data.py path/to/page.html

Guards: refuses a repeat (schedule_review1026 in the page) and a page without the v10.25 milestones line. The
workbooks are customer and branch documents and are not committed; their SHA-256 is checked and recorded.
"""
import copy, datetime, hashlib, json, os, re, sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE)); sys.path.insert(0, str(HERE.parent / 'toolchain'))
import baseplan1026 as B

BP_SHA = 'ba246d4af5f4faa436a1e208d034843343d13cb6bba033f811b9b966906e8771'
S9_SHA = '0a684fce53aa8ee93ccb9903e727e943244819c7142f58ddc7f234dd33ee4a26'
REFRE = re.compile(r'^\s*([A-Z]{1,3}\d{2,3})\b')

# joins settled by the branch's delivery dockets (Delivery Dockets sheet, 15 Oct 2026 export); each is a contract line
# the export did not touch, so the reader keeps its old join, and the old join is wrong on the ground
HAND_JOINS = {
    ('9968862', 93): ('P36', 'docket 26085104 delivered 1189410 on 17 Sep 2026 as "P36 Portable Building Shell"; the line names P36; the old join to HRP came from an earlier register'),
    ('9968862', 108): (None, 'the line names P67 and is Pending with 1189410, the number delivered to P36 on docket 26085104; P67 stands on site as unit 421107 on the record'),
    ('9968862', 64): ('P06', 'docket 26094121 delivered 1087453 on 30 Sep 2026 as "P06 Portable Building Shell"; the schedule wrote 1097453 for the same building'),
}


def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def record_matches(state):
    """Andrew's recorded numbers -> reference, from the shared record: typed asset numbers and the active unit records.
    A number recorded on two references is left out (the reader insists on one owner)."""
    docs = state['docs']; of = {}; dup = set()
    for k, v in (docs.get('assetNumbers') or {}).items():
        for n in (v.get('v') or []):
            n = str(n)
            if n in of and of[n] != k: dup.add(n)
            of.setdefault(n, k)
    for v in (docs.get('loads') or {}).values():
        if v.get('kind') == 'unit925-record' and not v.get('removedBy') and v.get('assetNo') and B.PLANT.fullmatch(str(v['assetNo'])):
            n = str(v['assetNo'])
            if n in of and of[n] != v['ref']: dup.add(n)
            of.setdefault(n, v['ref'])
    out = {}
    for n, k in of.items():
        if n in dup: continue
        out.setdefault(k, []).append(n)
    return out, sorted(dup)


# contracts the export carries that must not be priced beside what the page already prices for the same scope
HELD_OUT = {
    '9990776': 'Fencing. The page charges fencing from Advanced\'s measured dockets at the 2026 street card (Andrew, 19 Sep 2026); contract 9990776-STPS books the same scope at the same card rates with programme quantities (8,778 m clean, 2,775 m scrim, 2,350 m crowd control, 4,603 m demarcation), all Pending. Priced beside the dockets it would count the fencing twice ($398,991.88 at the page\'s rule). Held in the source history until the branch\'s lines and the dockets are reconciled.',
    '9968726': 'TV compound, superseded. The same four machines, 1321364, 1298390, 1220903 and 1300290, are booked again on 9992300-MEAD for the TV order of 14 Oct; the source audit flags each number on two priced lines. Mead to cancel 9968726; held in the source history so the compound is priced once.',
}

# the card the project manager chose, against the Rate 1 the export now carries. The page's own rule (v5.81): a VMS,
# forklift or water barrier charges by the day; a generator by its weekly rate over its hire (v10.04); everything
# else charges Rate 1 once as the whole-event bill. So Rate 1 is compared with the card figure of the same shape.
DAILY_KINDS = {'vms', 'forklift', 'barrier', 'trakmat', 'ground protection', 'access', 'forklift accessory'}
def rate_check(D, rows):
    card = {it['item']: it for it in D['rate_match']['items'] if it.get('rate')}
    out = []
    for r in rows:
        if r.get('charge_line') or not isinstance(r.get('rate_1'), (int, float)): continue
        it = card.get(r.get('register_type'))
        if not it: continue
        daily = it.get('rate'); weekly = it.get('weekly_rate'); hire = it.get('hire_rate')
        kind = r.get('kind') or ''
        if kind in ('generator', 'lighting', 'lighting tower'): basis, expect = 'card daily x 7 (weekly)', round(daily * 7, 2) if daily else None
        elif kind in DAILY_KINDS: basis, expect = 'card daily', round(daily, 2) if daily else None
        else: basis, expect = "card hire rate (the whole-event bill)", round(hire, 2) if hire else None
        if expect is None: continue
        diff = round(r['rate_1'] - expect, 2)
        agrees = abs(diff) <= max(0.05, 0.01 * expect)
        if not agrees and kind not in ('generator', 'lighting', 'lighting tower') and kind not in DAILY_KINDS:
            # a toilet or block is written for its hire term: a whole multiple of the card's weekly rate agrees too
            if weekly and abs(r['rate_1'] / weekly - round(r['rate_1'] / weekly)) < 0.01: agrees, basis = True, 'card weekly x %d weeks' % round(r['rate_1'] / weekly)
        out.append({'line': f"{r['rental_contract']}/{r['line']}", 'description': r.get('description'), 'register_type': r.get('register_type'), 'kind': kind, 'rate_1': r['rate_1'], 'rate_type': r.get('rate_type'), 'card': expect, 'basis': basis, 'card_line': it.get('rate_line'), 'difference': diff, 'agrees': agrees})
    return out


def refresh_baseplan(D, xlsx, state):
    assert sha(xlsx) == BP_SHA, 'wrong Baseplan workbook'
    R = D['rental_on_hire']; old = copy.deepcopy(R)
    oldrows = {(str(r['rental_contract']), int(r['line'])): r for r in old['rows']}
    NEW, TABS = B.read_export(xlsx)
    matches, dup = record_matches(state)
    known = {a['key'] for a in D['assets']} | {p['key'] for p in D['plant_lines']['lines']} | {u['task_id'] for u in D.get('unreferenced') or []}
    matches = {k: v for k, v in matches.items() if k in known}
    log = B.apply(D, NEW, TABS, matches, rejoin_record=False)
    log['record_numbers_with_two_owners_left_out'] = dup
    rows = {(str(r['rental_contract']), int(r['line'])): r for r in R['rows']}
    hand = []
    for k, (key, why) in HAND_JOINS.items():
        r = rows.get(k)
        if not r: continue
        was = r.get('match') or {}
        r['match'] = {'state': 'settled by delivery docket' if key else 'unmatched', 'to': 'asset' if key else None, 'key': key, 'task_id': None,
                      'via': 'delivery docket and the line\'s own reference' if key else None,
                      'basis': why + ' (set 15 Oct 2026 from the Delivery Dockets sheet of the same export)'}
        hand.append({'line': f'{k[0]}/{k[1]}', 'was': was.get('key') or was.get('task_id'), 'now': key, 'why': why})
    log['hand_joins'] = hand
    # assignments follow the joins that moved by hand
    AS = R['assignments']
    for key in sorted({h['was'] for h in hand if h['was']} | {h['now'] for h in hand if h['now']}):
        lines = [r for r in R['rows'] if r.get('match') and r['match'].get('to') == 'asset' and r['match'].get('key') == key]
        if lines:
            first = lines[0]; dl = next((r for r in lines if r.get('delivered')), None)
            AS[key] = {'branch_code': first['branch_code'], 'rental_contract': str(first['rental_contract']), 'via': first['match']['via'], 'basis': first['match']['basis'],
                       'lines': [f"{r['rental_contract']}/{r['line']}" for r in lines], 'all_contracts': sorted({str(r['rental_contract']) for r in lines}),
                       'delivered': {'date': dl['start_date'], 'delivery_number': dl['delivery_number']} if dl else None}
        elif key in AS and AS[key].get('via') in ('asset number', 'delivery docket', 'delivery docket and the line\'s own reference'):
            AS[key] = {'branch_code': AS[key]['branch_code'], 'rental_contract': AS[key]['rental_contract'], 'via': 'kind of item', 'basis': 'its numbered line now joins another reference', 'lines': [], 'all_contracts': [AS[key]['rental_contract']], 'delivered': None}
    for c in R['contracts']:
        rs = [r for r in R['rows'] if str(r['rental_contract']) == str(c['rental_contract'])]
        c['assets'] = sorted({r['match']['key'] for r in rs if r.get('match') and r['match'].get('to') == 'asset' and r['match'].get('key')})
        c['matched'] = sum(1 for r in rs if r.get('match') and r['match'].get('to'))
    removed = sorted(set(oldrows) - set(rows))
    hist = old.get('source_history949') or {'rows': []}
    held = []
    for r in R['rows']:
        why = HELD_OUT.get(str(r['rental_contract']))
        if why:
            r['held_out1026'] = why; held.append(copy.deepcopy(r))
    R['rows'] = [r for r in R['rows'] if str(r['rental_contract']) not in HELD_OUT]
    R['contracts'] = [c for c in R['contracts'] if str(c['rental_contract']) not in HELD_OUT]
    R['summary']['lines'] = len(R['rows']); R['summary']['contracts'] = [c for c in R['summary'].get('contracts', []) if str(c) not in HELD_OUT]
    R['summary']['held_out1026'] = {c: len([h for h in held if str(h['rental_contract']) == c]) for c in HELD_OUT}
    R['source_history949'] = {'source': old.get('source'), 'supplied_on': old.get('supplied_on'), 'superseded_by_sha256': BP_SHA,
                              'rows': [r for r in hist.get('rows') or []] + [copy.deepcopy(oldrows[k]) for k in removed] + held,
                              'held_out1026': {c: why for c, why in HELD_OUT.items()}}
    log['held_out'] = [{'line': f"{h['rental_contract']}/{h['line']}", 'description': h.get('description'), 'status': h.get('status_as_written'), 'why': HELD_OUT[str(h['rental_contract'])][:60]} for h in held]
    log['rate_check'] = rate_check(D, R['rows'])
    R['source_audit949'] = {'overlappingNumberGroups': len(log['flags']), 'note': old['source_audit949']['note']}
    R['record_id'] = 'baseplan-contracts-2026-10-15'; R['workbook'] = 'Baseplan_SuperCars_2026-10-15.xlsx'
    R['source'] = 'Baseplan SuperCars.xlsx — latest workbook supplied by Andrew Fisher, 15 Oct 2026 (13 contracts and the branch\'s Delivery Dockets sheet). Contract information; recorded site allocations remain authoritative.'
    R['supplied_by'] = 'Andrew Fisher'; R['supplied_on'] = '2026-10-15'; R['source_sha256'] = BP_SHA
    R['contracts_count1026'] = len(TABS)
    # the two supplier footers, read from the same cells as v9.49 (checked to be there, unchanged)
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, data_only=True, read_only=True); evidence = []
    for sheet, top, code0 in [('9974042-KINP', 11, 'ROY002'), ('9961265-STPS', 27, 'PRE808')]:
        ws = wb[sheet]; con, branch = sheet.split('-'); code = B.txt(ws.cell(top, 1).value)
        assert code == code0 and B.txt(ws.cell(top + 3, 1).value).lower() == 'cost rate', 'supplier note moved'
        name = B.txt(ws.cell(top + 1, 2).value); sa = B.txt(ws.cell(top + 2, 2).value); rate = ws.cell(top + 3, 2).value
        lines = [r['line'] for r in R['rows'] if r['rental_contract'] == con and r['supplier_sub_rental'] == code]
        prev = next((e for e in old.get('supplier_cost_evidence949') or [] if e['supplierCode'] == code), None)
        assert prev and prev['costRate'] == rate and prev['supplierName'] == name, 'supplier footer changed - read it before trusting it'
        e = dict(prev); e.update({'lines': lines, 'sourceSha256': BP_SHA}); evidence.append(e)
    wb.close(); R['supplier_cost_evidence949'] = evidence
    active = {str(r['rental_contract']) + '/' + str(r['line']) for r in R['rows']}
    for v in R['assignments'].values(): v['lines'] = [k for k in v.get('lines', []) if k in active]
    return log, len(NEW), len(TABS)


# ------------------------------------------------------------------------------------------------- Schedule 9
def read_schedule9(xlsx):
    import openpyxl
    assert sha(xlsx) == S9_SHA, 'wrong Schedule 9 workbook'
    wb = openpyxl.load_workbook(xlsx, data_only=True, read_only=True); out = {}
    for ws in wb.worksheets:
        if ws.title in ('BOQ', 'Labour'): continue
        rows = list(ws.iter_rows(values_only=True))
        if not rows: continue
        H = [str(h or '').strip().lower() for h in rows[0]]
        def col(name): return H.index(name) if name in H else None
        ci = {k: col(k) for k in ('date', 'product', 'item', 'qty', 'local', 'notes', 'asset', 'dd')}
        lst = []
        for i, r in enumerate(rows[1:], start=2):
            g = lambda k: (r[ci[k]] if ci[k] is not None and ci[k] < len(r) else None)
            d = g('date'); d = d.date().isoformat() if isinstance(d, datetime.datetime) else (d.isoformat() if isinstance(d, datetime.date) else None)
            local = str(g('local') or '').strip(); m = REFRE.match(local)
            lst.append({'row': i, 'date': d, 'product': str(g('product') or '').strip(), 'item': str(g('item') or '').strip(), 'qty': g('qty'), 'local': local, 'ref': m.group(1) if m else None, 'notes': str(g('notes') or '').strip(), 'asset': str(g('asset') or '').strip(), 'dd': str(g('dd') or '').strip()})
        out[ws.title] = lst
    wb.close(); return out


def norm(s): return re.sub(r'[^a-z0-9]', '', str(s or '').lower())


def apply_schedule9(D, xlsx, state):
    S9 = read_schedule9(xlsx); dl = state['docs'].get('delivery') or {}
    on_site = {k for k, d in dl.items() if d.get('done') or d.get('state') == 'on site'}
    A = {a['key']: a for a in D['assets']}
    # page events by (sheet, key, normalised item), in order
    idx = {}
    for a in D['assets']:
        for e in a.get('events') or []:
            idx.setdefault((e.get('sheet'), a['key'], norm(e.get('item'))), []).append((a, e))
    used = set(); dates_applied = []; dates_held = []; qtys = []; unmatched_rows = []
    for sheet, rows in S9.items():
        for r in rows:
            if not r['ref'] or r['ref'] not in A: continue
            cands = [x for x in idx.get((sheet, r['ref'], norm(r['item'])), []) if id(x[1]) not in used]
            if not cands:
                unmatched_rows.append({'sheet': sheet, 'row': r['row'], 'ref': r['ref'], 'item': r['item'], 'date': r['date'], 'qty': r['qty']}); continue
            a, e = cands[0]; used.add(id(e))
            # quantity: a whole number the sheet writes, different from the page
            q = r['qty']
            if isinstance(q, str) and re.fullmatch(r'\s*\d{1,4}\s*', q): q = int(q)   # the sheet writes some counts as text
            # a whole number under 10,000: the schedule once put an asset number in a quantity cell (R02)
            if isinstance(q, (int, float)) and not isinstance(q, bool) and float(q).is_integer() and 1 <= int(q) < 10000 and e.get('quantity_raw') != int(q):
                qtys.append({'sheet': sheet, 'row': r['row'], 'ref': r['ref'], 'item': e.get('item'), 'from': e.get('quantity_raw'), 'to': int(q)})
                e['quantity_correction'] = {'from': e.get('quantity_raw'), 'to': int(q), 'source': 'Schedule 9, 15 Oct 2026', 'cell': f"'{sheet}'!D{r['row']}"}
                e['quantity_raw'] = int(q); e['quantity_display'] = str(int(q))
            # date: moved by the schedule, unless the reference is already on site and this is its first day
            if r['date'] and e.get('date') and r['date'] != e['date']:
                first = min(x['date'] for x in a['events'] if x.get('date'))
                entry = {'sheet': sheet, 'row': r['row'], 'ref': r['ref'], 'item': e.get('item'), 'from': e['date'], 'to': r['date']}
                if r['ref'] in on_site and e['date'] == first:
                    entry['held'] = 'on site on the record; its first day is history'; dates_held.append(entry)
                else:
                    e['date_correction'] = {'from': e['date'], 'to': r['date'], 'source': 'Schedule 9, 15 Oct 2026', 'cell': f"'{sheet}'!A{r['row']}"}
                    e['date'] = r['date']; e['date_display'] = datetime.date.fromisoformat(r['date']).strftime('%-d %b')
                    dates_applied.append(entry)
    # new rows for references the page has: the two fridges
    added = []
    for sheet, rows in S9.items():
        for r in rows:
            if not r['ref'] or r['ref'] not in A: continue
            if r['product'].lower().startswith('furniture') and r['item'].lower() == 'fridge' and not any(e for e in A[r['ref']].get('events') or [] if e.get('sheet') == sheet and norm(e.get('item')) in ('fridgelge', 'fridge')):
                a = A[r['ref']]
                e = {'task_id': None, 'sheet': sheet, 'phase': 'Build', 'date': r['date'], 'date_display': datetime.date.fromisoformat(r['date']).strftime('%-d %b'), 'activity': None, 'item': 'Fridge Lge',
                     'date_as_written': None, 'date_correction': None, 'quantity_display': '1', 'quantity_flag': None, 'quantity_raw': 1, 'quantity_correction': None,
                     'source_range': f"'{sheet}'!A{r['row']}:I{r['row']}", 'dd': r['dd'] or None, 'carrier': None, 'load_time': None, 'transport_cost_text': None, 'transport_cost': None,
                     'note': 'Schedule 9 row, 15 Oct 2026', 'movement': 'place', 'movement_basis': 'inferred from the Build phase — no activity word in the source', 'movement_stated': False}
                a.setdefault('events', []).append(e); added.append({'sheet': sheet, 'row': r['row'], 'ref': r['ref'], 'item': 'Fridge Lge', 'date': r['date']})
                for w in D['weeks']:
                    if w['sheet'] == sheet: w['rows'] = int(w['rows']) + 1
    # first and last day follow the events (the rule every asset on the page already keeps)
    touched = {x['ref'] for x in dates_applied} | {x['ref'] for x in added}
    for k in touched:
        ds = [e['date'] for e in A[k].get('events') or [] if e.get('date')]
        if ds: A[k]['first_date'] = min(ds); A[k]['last_date'] = max(ds)
    held_rows = [{'sheet': s, 'row': r['row'], 'local': r['local'], 'item': r['item'], 'qty': r['qty'], 'date': r['date']} for s, rows in S9.items() for r in rows if r['ref'] and r['ref'] not in A and r['product'].upper().startswith('VMS')]
    return {'dates_applied': dates_applied, 'dates_held': dates_held, 'quantities': qtys, 'rows_added': added, 'rows_without_a_page_event': unmatched_rows, 'vms_rows_held': held_rows}


# the schedule's copies of a number that belongs to another reference: the register shows them both until somebody
# says which, and the page's own rule hides the number on both. Each is settled by the record and the 15 Oct export.
SCHEDULE_COPIES = [
    ('P63', '1282487', 'P63 stands as unit 960719 on the record and on Baseplan line 121; 1282487 is HRP (Schedule 9, Week 5 row for the Hire Republic crib room)'),
    ('P27', '1097343', 'docket 26076747 delivered 1097343 to P18 on 16 Sep; P27 is not on site'),
    ('P53', '1327222', 'docket 26126001 delivered 1327222 to P66 on 14 Oct'),
    ('P36', '1327222', 'docket 26126001 delivered 1327222 to P66 on 14 Oct; P36 is 1189410 (docket 26085104)'),
    ('P46', '1327213', 'docket 26064087 delivered 1327213 to P04 on 15 Sep; P46 is 1327215 on the record'),
    ('P67', '1189410', 'docket 26085104 delivered 1189410 to P36 on 17 Sep; P67 is the commentary booth, unit 421107 on the record'),
    ('GN06', '1283681', 'the 15 Oct export carries 1283681 as a spare generator; GN06 is 1320828, confirmed on site 15 Oct'),
]


def settle_schedule_copies(D):
    A = {a['key']: a for a in D['assets']}; out = []
    for key, no, why in SCHEDULE_COPIES:
        a = A.get(key)
        if not a or no not in (a.get('asset_numbers') or []): continue
        a['asset_numbers'] = [n for n in a['asset_numbers'] if n != no]
        a.setdefault('asset_notes', []).append({'kind': 'number settled', 'number': no, 'note': 'the schedule wrote ' + no + ' here; ' + why + ' (settled 15 Oct 2026)'})
        a['asset_no_state'] = a.get('asset_no_state') if a['asset_numbers'] else 'settled: the schedule\'s number belongs to another reference'
        out.append({'ref': key, 'number_removed': no, 'why': why})
    return out


def main():
    a = sys.argv[1:]
    if not a: sys.exit(__doc__)
    page = Path(a[0]); html = page.read_text(encoding='utf-8')
    if 'schedule_review1026' in html: sys.exit('patch_v1026_data refuses a repeat')
    if 'ms1025-script' not in html: sys.exit('patch_v1026_data: not the v10.25 line - wrong base')
    bp = os.environ['BASEPLAN1026_XLSX']; s9 = os.environ['SCHEDULE9_XLSX']; state = json.load(open(os.environ['RECORD_JSON']))
    at = html.index('const DATA = ') + len('const DATA = ')
    D, n = json.JSONDecoder().raw_decode(html[at:])
    assert json.dumps(D, ensure_ascii=False, separators=(',', ':')) == html[at:at + n], 'DATA does not round-trip byte for byte'
    before = copy.deepcopy(D)
    copies = settle_schedule_copies(D)
    bplog, nlines, ntabs = refresh_baseplan(D, bp, state)
    s9log = apply_schedule9(D, s9, state)
    D['baseplan_review1026'] = {'schema': 1, 'author': 'Andrew Fisher', 'reviewed': '2026-10-15', 'source': 'Baseplan SuperCars.xlsx, exported 15 Oct 2026 (uploaded as 892b9fab-Baseplan_SuperCars.xlsx)', 'sha256': BP_SHA,
                                'contract_tabs': ntabs, 'lines': nlines, 'record_version': state.get('version'),
                                'added': bplog['added'], 'removed_to_history': bplog['removed'], 'changed': len(bplog['changed']), 'changed_lines': bplog['changed'],
                                'rejoined_by_filled_number': bplog['rejoined'], 'held': bplog['held'], 'record_would_move': bplog['record_would_move'], 'hand_joins': bplog['hand_joins'],
                                'repeated_numbers': bplog['flags'], 'plant_line_copies': bplog['plant_line_copies'], 'schedule_copies_settled': copies,
                                'held_out': bplog['held_out'], 'held_out_why': HELD_OUT, 'rate_check': bplog['rate_check'],
                                'rule': 'Contract information as exported; the site record stays authoritative for what stands where. A line the export no longer has is kept in the source history, out of the active totals.'}
    D['schedule_review1026'] = {'schema': 1, 'author': 'Andrew Fisher', 'reviewed': '2026-10-15', 'source': 'GC500 26 Schedule 9.xlsx (uploaded as 85b798de-GC500_26_Schedule_9.xlsx)', 'sha256': S9_SHA, 'comparison': 'the Schedule 7 events the page held',
                                'applied': s9log, 'rule': 'A quantity or a day the schedule changed moves on the page; a reference already on site keeps its first day; VMS pick-up rows stay with the VMS plant line (no page reference).'}
    # nothing else moved
    keep = {'rental_on_hire', 'plant_lines', 'assets', 'weeks', 'baseplan_review1026', 'schedule_review1026'}
    assert {k: v for k, v in D.items() if k not in keep} == {k: v for k, v in before.items() if k not in keep}, 'data outside the plan changed'
    for x, y in zip(D['assets'], before['assets']):
        assert x['key'] == y['key']
        assert {k: v for k, v in x.items() if k not in ('events', 'first_date', 'last_date', 'asset_numbers', 'asset_notes', 'asset_no_state')} == {k: v for k, v in y.items() if k not in ('events', 'first_date', 'last_date', 'asset_numbers', 'asset_notes', 'asset_no_state')}, 'asset ' + x['key'] + ' changed outside its events and settled numbers'
    html = html[:at] + json.dumps(D, ensure_ascii=False, separators=(',', ':')) + html[at + n:]
    page.write_text(html, encoding='utf-8')
    rep = {'baseplan': {k: (len(v) if isinstance(v, list) else v) for k, v in bplog.items()}, 'schedule9': {k: len(v) for k, v in s9log.items()}}
    (HERE / 'data_log_1026.json').write_text(json.dumps({'baseplan': bplog, 'schedule9': s9log, 'schedule_copies': copies}, indent=1, ensure_ascii=False) + '\n')
    print('patch_v1026_data applied:', json.dumps(rep))


if __name__ == '__main__':
    main()
