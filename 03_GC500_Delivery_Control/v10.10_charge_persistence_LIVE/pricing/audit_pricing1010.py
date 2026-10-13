"""Author: Andrew Fisher. Read-only source-rate and quantity audit; snapshots stay private."""
import collections
import datetime
import hashlib
import json
import os
import re
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from openpyxl import load_workbook

snapshot_path=Path(os.environ['SNAPSHOT'])
card_path=Path(os.environ['RATE_CARD'])
s=json.loads(snapshot_path.read_text())
w=load_workbook(card_path,data_only=True,read_only=True)
sheet=w['Street Rate Card 2026']
norm=lambda v: re.sub(r'\s+',' ',str(v or '').strip().lower())
grid={norm(row[0].value):(row[0].row,[c.value for c in row]) for row in sheet.iter_rows() if row[0].value}
close=lambda a,b: isinstance(a,(int,float)) and isinstance(b,(int,float)) and abs(a-b)<0.000001
cents=lambda v: float(Decimal(str(v)).quantize(Decimal('.01'),rounding=ROUND_HALF_UP))
checks=[]
def check(name,passed,detail=None):
    checks.append(dict(name=name,pass_=bool(passed),**({'detail':detail} if detail is not None else {})))

seen=set();source_work=[]
for item in s['items']:
    card=item['card'];line=norm(card.get('line'));key=(line,json.dumps(card.get('labour_per_piece'),sort_keys=True))
    if key in seen or line not in grid:continue
    seen.add(key);source_row,cells=grid[line]
    for work in (card.get('labour_per_piece') or {}).get('lines',[]):
        if work.get('money')!='rate':continue
        col={'install':4,'steps':5,'levelling':6,'cleaning':7,'fire_ext':8,'demob':9}.get(work['key'])
        if col is None:continue
        if 66<=source_row<=79 and work['key'] in ('install','demob'):
            expected=cells[1]/2
        elif source_row<=39:
            expected=cells[col]
        else:continue
        passed=close(work['rate'],expected)
        check('Work source '+line+' '+work['key'],passed)
        source_work.append(dict(item=line,charge=work['key'],row=source_row,expected=expected,actual=work['rate'],matches=passed))

for item in s['items']:
    if item['received']<=0:continue
    for row in item['charges']['rows']:
        if row['qty'] is None:continue
        check('Within approved quantity '+item['ref']+' '+item['item']+' '+row['key'],row['qty']<=item['quantity'])
        if row['rate'] is not None:check('Service multiplication '+item['ref']+' '+item['item']+' '+row['key'],close(row['value'],cents(row['qty']*row['rate'])))

generator_grid={int(m.group(1)):(source_row,cells) for key,(source_row,cells)in grid.items() if (m:=re.search(r'generator\s*(\d+)\s*kva',key)) and 66<=source_row<=79}
for row in s['generators']['rows']:
    source_row,cells=generator_grid[row['pricedKva']]
    if row['rateState']=='card':
        check('Generator published daily '+str(row['ref'] or row['id']),close(row['dailyRate'],cells[2]))
        check('Generator published standdown '+str(row['ref'] or row['id']),close(row['standdownRate'],cells[3]))
    start=row['deliveredDate'] if row['actualIsConfirmed'] else row['plannedDate']
    if start:
        day=datetime.date.fromisoformat(start);friday=datetime.date(2026,10,23);sunday=datetime.date(2026,10,25)
        standby=max(0,(friday-day).days);race=max(0,(sunday-max(day,friday)).days+1)
        check('Generator delivery-inclusive standdown '+str(row['ref'] or row['id']),row['standdownDays']==standby)
        check('Generator Friday to Sunday only '+str(row['ref'] or row['id']),row['eventDays']==race and row['offhire']=='2026-10-25')

ids=[str(r['rental_contract'])+'|'+str(r['line']) for r in s['contracts']]
check('No duplicate contract-line identity in job totals',len(ids)==len(set(ids)))
slots=[(r['ref'],r['disc'],r['item'],r['key'],r['state'])for r in s['plan']['slots']]
check('No duplicate service slot in actual or forecast',len(slots)==len(set(slots)))
check('Known onsite subtotal equals sum of displayed item totals',close(s['bill']['known'],cents(sum(r['total'] for r in s['bill']['rows']))))
check('No unauthorised damage waiver is priced',all(r.get('now') is None and r.get('job') is None for r in s['pl']['rev'] if str(r.get('code'))=='1015'))
check('All arrived item groups have a branch',all((i.get('branch')or{}).get('code') for i in s['items'] if i['received']>0))
check('All existing integration controls pass',all(c['pass'] for c in s['checks']))

overallocated=[]
for i in s['items']:
    rows=(i.get('hire')or{}).get('lines',[])
    if rows and all(not r.get('assetNo') for r in rows) and sum(r.get('quantity')or 0 for r in rows)>i['quantity']:
        overallocated.append(dict(ref=i['ref'],item=i['item'],approved=i['quantity'],contractQuantity=sum(r['quantity']or 0 for r in rows),contracts=[r['id']for r in rows]))
serials=collections.defaultdict(list)
for r in s['contracts']:
    if r['asset_no_is_plant_number'] and not r['charge_line']:serials[r['asset_no']].append(r)
duplicate_review=[dict(asset=asset,contracts=[dict(contract=r['rental_contract'],line=r['line'],status=r['status_as_written'],start=r['start_date']or r['booked_delivery_date'],end=r['term_date']or r['expected_term_date'])for r in rows])for asset,rows in serials.items() if len(rows)>1]
held=[r for r in s['transport']['rows'] if r['state']=='held' and (r.get('receivedQuantity1006')or 0)>0]
result=dict(author='Andrew Fisher',sha256=s['sha256'],nativeVersion=s['nativeVersion'],cardSha256=hashlib.sha256(card_path.read_bytes()).hexdigest(),itemGroups=len(s['items']),arrivedGroups=sum(i['received']>0 for i in s['items']),checks=len(checks),failures=[c for c in checks if not c['pass_']],workSourceComparisons=len(source_work),overallocated=overallocated,repeatedSerialReview=duplicate_review,arrivedTransportHolds=len(held),arrivedTransportReferences=len({r['ref']for r in held}),unresolvedPackageHolds=len(s['bill']['holds']),rateArithmetic='Source arithmetic verified for mapped work and generator rates; unresolved scope, duplicate allocations and missing inputs are not certified complete.')
out=Path(os.environ.get('AUDIT_OUT',str(snapshot_path.parent/'source-audit.json')))
out.write_text(json.dumps(result,indent=2)+'\n')
print(json.dumps({k:v for k,v in result.items()if k!='repeatedSerialReview'},indent=2))
if result['failures']:raise SystemExit(1)
