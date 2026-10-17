#!/usr/bin/env python3
"""Read source workbooks and existing snapshot; write review artefacts only.

Author: Andrew Fisher. No operational records are changed.
"""
import argparse
import hashlib
import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def value(v):
    return v.isoformat() if isinstance(v, (date, datetime)) else v


def read(path):
    book = load_workbook(path, read_only=True, data_only=False)
    rows = {s.title: list(s.values) for s in book}
    meta = {
        'filename': path.name, 'bytes': path.stat().st_size,
        'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
        'created': value(book.properties.created),
        'modified': value(book.properties.modified),
        'sheets': [{'name': n, 'rows': len(r), 'columns': max(map(len, r), default=0)}
                   for n, r in rows.items()],
    }
    book.close()
    return rows, meta


def changes(before, after):
    out = []
    for name in sorted(set(before) | set(after)):
        a, b = before.get(name, []), after.get(name, [])
        for row in range(max(len(a), len(b))):
            ar, br = a[row] if row < len(a) else (), b[row] if row < len(b) else ()
            for col in range(max(len(ar), len(br))):
                av, bv = ar[col] if col < len(ar) else None, br[col] if col < len(br) else None
                if av != bv:
                    out.append({'sheet': name, 'cell': f'{get_column_letter(col+1)}{row+1}',
                                'before': value(av), 'after': value(bv)})
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--attachments', type=Path, required=True)
    ap.add_argument('--snapshot', type=Path, required=True)
    ap.add_argument('--output', type=Path, required=True)
    ap.add_argument('--private-evidence', type=Path, required=True)
    a = ap.parse_args()
    files = {
        'programme': a.attachments / '15e4fda3-c153-4a4f-9998-7b948008e797/1-P003-26003-01-Programme-Coates-Inf-V1-14.8.26.xlsx',
        'prior_programme': a.attachments / '97e3ae1c-6e01-4dd1-b6c2-807968018afe/4-P003-26003-01-Programme-Coates-Inf-V1-14.8.26-2-.xlsx',
        'schedule6': a.attachments / '15e4fda3-c153-4a4f-9998-7b948008e797/2-GC500-26-Schedule-6-.xlsx',
        'schedule8': a.attachments / 'a60922c4-c761-44fe-b9e9-14ccb8970fc1/1-GC500-26-Schedule-8-.xlsx',
    }
    loaded = {k: read(p) for k, p in files.items()}
    p_diff = changes(loaded['prior_programme'][0], loaded['programme'][0])
    s_diff = changes(loaded['schedule6'][0], loaded['schedule8'][0])
    assert [(x['sheet'], x['cell'], x['before'], x['after']) for x in p_diff] == [
        ('CON WK 2', 'I101', None, 'Complete'),
        ('CON WK 2', 'I102', None, 'Complete'),
        ('CON WK 2', 'I103', None, 'Complete'),
    ], 'Programme source has changed: review again before producing report.'
    assert loaded['schedule6'][1]['modified'] < loaded['schedule8'][1]['modified']
    snapshot = json.loads(a.snapshot.read_text())
    items = snapshot['items']

    def item(ref):
        matches = [x for x in items if x.get('ref') == ref]
        assert len(matches) == 1, (ref, len(matches))
        return matches[0]

    def fleet(ref):
        return ', '.join(str(x.get('assetNo')) for x in item(ref).get('physicalUnits', [])
                         if x.get('assetNo')) or 'No confirmed asset'

    flag_evidence = []
    for ref, cell in [('GN13', 'I101'), ('P69', 'I102'), ('P45', 'I103')]:
        record = item(ref)
        assert record['received'] == 1 and record['delivery']['state'] == 'on site'
        assert record['delivery']['done'] is True
        flag_evidence.append({'ref': ref, 'cell': cell, 'received': 1,
                              'on_site': True, 'complete': True, 'actual_quantity_change': 0})

    unchanged = [s for s in loaded['schedule6'][0] if not any(x['sheet'] == s for x in s_diff)]
    assert all(s in unchanged for s in ['Demob Week 1', 'Demob Week 2', 'Demob Week 3'])
    evidence = {
        'author': 'Andrew Fisher', 'purpose': 'Read-only source comparison; no operational import',
        'sources': {k: m for k, (_, m) in loaded.items()},
        'snapshot': {'at': snapshot['at'], 'nativeVersion': snapshot['nativeVersion'],
                     'file_sha256': hashlib.sha256(a.snapshot.read_bytes()).hexdigest(),
                     'parent_sha256': snapshot['sha256']},
        'programme_changes': p_diff, 'programme_flags_already_recorded': flag_evidence,
        'schedule_unchanged_sheets': unchanged,
        'schedule_cell_change_count': len(s_diff),
        'rule': 'Andrew Fisher\u2019s confirmed onsite records take priority over planning documents.',
    }
    here = Path(__file__).resolve().parent
    (here / 'source-review.json').write_text(json.dumps(evidence, indent=2) + '\n')
    a.private_evidence.mkdir(parents=True, exist_ok=True)
    (a.private_evidence / 'full-cell-comparison.json').write_text(json.dumps({
        'author': 'Andrew Fisher', 'programme': p_diff, 'schedule6_to_8': s_diff,
    }, indent=2) + '\n')

    # These are review findings, not new records, forecasts or customer charges.
    existing = []
    for ref, cell in [('GN13', 'I101'), ('P69', 'I102'), ('P45', 'I103')]:
        existing.append([
            ref, f"1 × {item(ref)['item']} · asset {fleet(ref)}\nOn site and complete: 12 Oct 2026",
            f"NEW flag only: Programme CON WK 2!{cell} now says Complete.",
            'Already recorded. No extra arrival, quantity or charge from this flag.',
        ])
    existing += [
        ['File revision', 'Schedule 8 is newer: modified 12 Oct.\nSupplied Schedule 6: modified 30 Sep.',
         'Programme filename says August; only three cell values differ from the previous Programme.',
         'Keep Schedule 8 as the newer planning source. Keep confirmed site records first.'],
        ['Demob / rates', 'All three Schedule demob tabs are unchanged.\nBOQ and Labour tabs are also unchanged.',
         'Schedule 6 compared with Schedule 8: exact cell-value comparison.',
         'No new pickup plan or rates to import from these files.'],
        ['GN13 / GN18', f"Current fleet: GN13 {fleet('GN13')}; GN18 {fleet('GN18')}.\nBoth received: 1.",
         'Schedule 8 changes earlier planned fleet numbers; the current records already match.',
         'Keep the current assets. Do not add a second unit.'],
        ['GN04 / GN25', f"Current allocation: GN04 {fleet('GN04')}; GN25 {fleet('GN25')}.\nBoth received: 0.",
         'Schedule 8 updates the planned allocation; the current records already match.',
         'Keep as planned. An allocated number is not a new arrival confirmation.'],
        ['WC09', 'Received: 4 × FWF + 6 × Pee Panel + 2 × 6 m toilet block.',
         'Both schedules list these three item groups. Programme F99 calls the smaller units 10 single toilets.',
         'Keep the item groups separate. Do not replace them with 10 FWF.'],
        ['WAU / subhire', 'WAU: 1 fridge + 2 portable air conditioners are already recorded as source requests.',
         'Schedule 8 Event Week rows 46–47. Added “event p” notes also support Event Portables ownership.',
         'Keep the existing requests and subhire classification. No duplicate assets or costs.'],
    ]
    fixes = [
        ['P46 / P51', 'Source conflict',
         f"Confirmed: P46 {fleet('P46')}; P51 {fleet('P51')}. Both received 1.\nSchedules allocate these differently.",
         'Correct the planning documents to the confirmed site asset numbers. Keep actuals unchanged.',
         'Schedules 6 & 8: Week 2!I6 and I8'],
        ['GN10', 'Planning change',
         'Schedule 8 changes 20 kVA to 45 kVA. Current requested item remains 20 kVA; received 0.',
         'Resolve the requested size before changing the hire rate or asset allocation.',
         'Schedule 8: Event Week!C4, F4; Schedule 6: C4'],
        ['GN19 / GN24', 'Allocation notes',
         'GN19 requests 60 kVA but the allocated fleet note says 80 kVA. GN24 fleet is still provisional.',
         'Keep GN19 requested-size pricing; confirm GN24 fleet before adding an asset number.',
         'Schedule 8: Event Week!C5, I5:I6'],
        ['P25', 'Source conflict',
         'Schedule 6 and Programme say 4.8 m; Schedule 8 says 6 m. Same asset number; received 0.',
         'Confirm the building size against the asset/baseplan contract before changing the description or price.',
         'Schedule 6/8: Week 1!C15; Programme: B107'],
        ['WB06', 'Planning date',
         '154 × TL2: Schedule 8 moves the plan from 16 Oct to 15 Oct. Embedded plan still shows 16 Oct; received 0.',
         'Check the effective daily run against 15 Oct. Do not change any confirmed arrival.',
         'Schedule 6: Week 1 row 36; Schedule 8: row 30'],
        ['VMS / Marriott / Tickford', 'Planned scope to link',
         'Schedule 8: event VMS 10 → 12; Marriott crossing 4 × TL2 on 18 Oct; Tickford 1 × FWF on 21 Oct.',
         'Link each planned row to its existing reference and contract before adding anything. Do not mark delivered.',
         'Schedule 8: Event Week rows 28, 40–42; Week 1 row 38'],
        ['Event forklifts / fencing loads', 'Planning reconciliation',
         'Schedule 8 changes forklift supply rows and adds fencing truck movements.',
         'Match existing forklift assets and delivery jobs first. Fencing truckload quantities are not fence-panel quantities.',
         'Schedule 8: Week 1 rows 8–9; Week 3 rows 18, 27–28; Week 2 row 10'],
        ['WC57 / WC59 / WC67', 'Older known conflict',
         'Programme swaps WC57/WC59 counts. Current WC59: 7 received; WC57: 0 of 2. WC67: 2 received; schedule calls for 4.',
         'Keep confirmed counts. Correct WC57/WC59 in Programme; check WC67 as 2 existing plus 2 planned, not 4 delivered.',
         'Programme Base information (4): F73:G73, F77, F104'],
        ['Programme dates / demob copy', 'Older known error',
         'P13–P16 dates show 2029; P69/P45 end dates show 2025. Final “DEMOB WK 1” tab contains September delivery rows.',
         'Correct source year errors and identify the real demob plan. Do not use the copied delivery tab for pickup dates.',
         'Programme Base information (4): D40:E43, E102:E103; final DEMOB WK 1'],
    ]

    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = 'Andrew Fisher'
    wb.properties.title = 'GC500 Programme and Schedule Review'
    wb.properties.subject = 'Existing truth and source items to fix; read-only comparison'
    navy, orange, pale, grey = '172D38', 'F47721', 'EAF0F3', '54636C'
    def sheet(name, columns, rows, widths, subtitle):
        ws = wb.create_sheet(name)
        n = len(columns)
        ws.sheet_view.showGridLines = False
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
        ws.cell(1, 1, f'GC500 2026 | {name}').font = Font(name='Aptos Display', size=20, bold=True, color='FFFFFF')
        ws.cell(1, 1).fill = PatternFill('solid', fgColor=navy)
        ws.row_dimensions[1].height = 36
        for row, text in [(2, subtitle), (3, 'Author: Andrew Fisher | Review: 14 Oct 2026 | Confirmed site records take priority')]:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n)
            ws.cell(row, 1, text).font = Font(name='Aptos', size=11, bold=row == 2, color=navy if row == 2 else grey)
            ws.cell(row, 1).alignment = Alignment(wrap_text=True, vertical='center')
            ws.row_dimensions[row].height = 30 if row == 2 else 24
        ws.append([None] * n)
        ws.append(columns)
        for c in ws[5]:
            c.font = Font(name='Aptos', bold=True, color='FFFFFF', size=11)
            c.fill = PatternFill('solid', fgColor=navy)
            c.alignment = Alignment(vertical='center', wrap_text=True)
            c.border = Border(bottom=Side(style='medium', color=orange))
        ws.row_dimensions[5].height = 28
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows(min_row=6, max_row=5 + len(rows)):
            ws.row_dimensions[row[0].row].height = 70 if name == 'Existing' else 94
            for c in row:
                c.font = Font(name='Aptos', size=11, color=navy, bold=c.column == 1)
                c.fill = PatternFill('solid', fgColor=pale if c.row % 2 == 0 else 'FFFFFF')
                c.alignment = Alignment(vertical='top', wrap_text=True)
                c.border = Border(bottom=Side(style='thin', color='D1DCE1'))
        footer = ws.max_row + 2
        source_lines = [
            'Sources: Programme = P003-26003-01 Programme Coates Inf V1 14.8.26; supplied file compared with the prior reviewed Programme.',
            'Schedule 6 = GC500 26 Schedule (6); Schedule 8 = GC500 26 Schedule (8). Exact original filenames and hashes are in the review evidence.',
            'Current record check: snapshot version 5796, captured 14 Oct 2026 AEST. This review makes no operational changes or cost-completeness claim.',
        ]
        for text in source_lines:
            ws.merge_cells(start_row=footer, start_column=1, end_row=footer, end_column=n)
            c = ws.cell(footer, 1, text)
            c.font = Font(name='Aptos', size=9, color=grey)
            c.alignment = Alignment(wrap_text=True, vertical='top')
            ws.row_dimensions[footer].height = 26
            footer += 1
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = 'B6'
        ws.auto_filter.ref = f'A5:{get_column_letter(n)}{5 + len(rows)}'
        ws.print_title_rows = '1:5'
        ws.print_area = f'A1:{get_column_letter(n)}{ws.max_row}'
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.oddFooter.left.text = 'Andrew Fisher | GC500 2026'
        ws.oddFooter.right.text = 'Page &P of &N'
        return ws
    sheet('Existing', ['Reference', 'Existing record', 'What the files add', 'Keep / action'], existing,
          [22, 49, 48, 49], 'Three new Programme completion flags. All three were already on site. No actual quantity changes.')
    sheet('Needs fixing', ['Reference', 'Type', 'What differs', 'Fix / keep now', 'Source cells'], fixes,
          [20, 20, 47, 46, 33], 'Planning and document corrections only. Older known errors are labelled. Confirmed site data stays unchanged.')
    a.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(a.output)
    check = load_workbook(a.output, read_only=True, data_only=False)
    assert check.sheetnames == ['Existing', 'Needs fixing']
    assert check.properties.creator == 'Andrew Fisher'
    assert not any(c.data_type == 'f' for ws in check for row in ws for c in row)
    check.close()
    print(json.dumps({'workbook': str(a.output), 'existing_rows': len(existing), 'needs_fixing_rows': len(fixes),
                      'programme_changes': len(p_diff), 'unchanged_schedule_tabs': unchanged,
                      'operational_writes': 0}, indent=2))


if __name__ == '__main__':
    main()
